# In Google Colab, run first: !pip install -q yfinance openpyxl scikit-learn

"""BeatOdds for Google Colab. Paste everything into ONE cell and run (about 10 to 14 minutes, calculation included).

Question: Polymarket lists "Will [company] beat quarterly earnings?" contracts. Do their prices tell investors more than the
analyst consensus and the companies' own beat history, and are they linked to the stock's reaction to the report?

Downloads (all free, no key): Polymarket Gamma and CLOB APIs (earnings and Fed contracts, hourly prices), Kalshi Trade API v2
(company KPI and Fed contracts), Yahoo Finance via yfinance (prices, analyst EPS estimates, option chains, Treasury yield),
FRED (2-year Treasury yield, when reachable). Optional, switched off: SEC EDGAR release times and FINRA short interest.

Workbook:
- Results A to D: accuracy, calibration, stock reaction, the priced-in table.
- Results E to L: run-up and drift, liquidity tiers, combined forecast, change in the odds, GAAP split, move size, disagreement.
- Results O: simple trading rules on the prediction market. P: Fed weeks and Treasury yields.
- Results Q: robustness (beta-adjusted returns, 3-day window, odds read earlier, volatility thirds, analyst-history model).
- Results R: the confident-miss gap at cut-offs from 0.70 to 0.90, with 95% intervals.
- Company sheet: pick a ticker. Live sheet: scenario card per upcoming report and odds for the next Fed meeting.
- Fed sheets: how well Kalshi and Polymarket priced each Fed decision since 2024.
Python only downloads and arranges the data (and fits the analyst-history model). Every statistic is a visible Excel formula.
Declare: Claude drafted this script; the group reviewed and ran it.
"""

import json
import re
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone

import numpy as np
import pandas as pd
import requests

# =============================================================================
# CONFIG
# =============================================================================
GROUP_NUMBER = "A2"
OUTPUT_FILE = f"Group{GROUP_NUMBER}_P1_Workbook.xlsx"
RAW_ZIP = f"Group{GROUP_NUMBER}_P1_RawData.zip"
START_DATE = "2025-01-01"        # earnings reports from this date
CUTOFF_HOUR_UTC = 10             # odds are read at 10:00 UTC on the report day (06:00 New York), before any announcement
LEAD_DAYS = 7                    # second reading of the odds, this many days earlier
ODDS_WINDOW_DAYS = 10            # hourly odds kept for the last N days before the cut-off
MIN_VOLUME = 1000                # USD traded, to drop empty markets
MIN_PRIOR_EVENTS = 20            # events needed before a running base rate is used
MIN_HISTORY = 4                  # past reports needed for a company's own beat rate
THREADS = 8
MARKET = "SPY"                   # market return used for abnormal returns
GAMMA = "https://gamma-api.polymarket.com"
CLOB = "https://clob.polymarket.com"
KALSHI = "https://api.elections.kalshi.com/trade-api/v2"
USE_KALSHI = True                # set to False to run Polymarket only
USE_OPTIONS = True               # option-implied move for upcoming reports (Live sheet)
DRIFT_DAYS = 10                  # trading days for the pre-report run-up and the post-report drift
LIQ_VOLUME = 10000               # USD volume for the liquid-only robustness check (editable on Assumptions)
LIQ_VOLUME_2 = 25000             # stricter liquidity tier (editable on Assumptions)
DP_BAND = 0.05                   # change in p over the last week counted as "rose" or "fell"
HIST_START = "2021-01-01"        # longer price history, upcoming-report companies only, for their past earnings moves
MIN_PAST_MOVES = 6               # past reactions needed before the options reading is shown
HIGH_P = 0.8                     # "confident crowd" threshold used in the priced-in tests and the scenario card
VOL_DAYS = 20                    # trading days for the volatility before the report
GAP = 0.2                        # crowd vs company record gap counted as a disagreement
USE_EDGAR = False                # SEC EDGAR: exact time of each earnings press-release filing (8-K item 2.02)
SEC_USER_AGENT = "IUM student project your.name@example.com"   # SEC asks for a name and contact email: put yours here
USE_FINRA = False                # FINRA: short interest twice a month (public, no key)
SI_LAG_DAYS = 12                 # short interest is published about 8 business days after its settlement date
TRADE_COST = 0.02                # assumed spread and fees per $1 prediction-market contract (Results part O)
USE_FED = True                   # Fed meetings: Kalshi and Polymarket odds, FRED Treasury yields (Fed sheets, Results part P)
FED_START = "2024-01-01"         # Fed meetings from this date
RECALC_IN_COLAB = True           # calculate every formula before downloading (installs LibreOffice in Colab, about 2 minutes)
FED_SERIES = ("KXFEDDECISION", "FEDDECISION")
FED_TAGS = ("fed-rates", "fed", "fomc", "federal-reserve")
FRED_CSV = "https://fred.stlouisfed.org/graph/fredgraph.csv"
FED_WEEK_DAYS = 3                # an earnings report within this many days of a Fed decision counts as a "Fed week"
YIELD_MOVE = 0.10                # 2-year yield change (percentage points, 20 trading days) counted as rising or falling
K_MIN_VOLUME = 200               # Kalshi contracts traded, to drop empty KPI markets
K_MIN_PRIOR_EVENTS = 10          # earlier Kalshi events needed for the running base rate
# Kalshi series tickers that do not contain the stock ticker. None = no listed stock, skip.
K_OVERRIDES = {"KXWALMARTA": "WMT", "KXUBERTRIPSA": "UBER", "KXCOINBASE": "COIN", "KXKLARNAA": "KLAR", "KXMETA": "META",
               "KXDIS": "DIS", "KXDISA": "DIS", "KXTSM": "TSM", "KXSPCXA": None, "GOOGLESHARE": None, "OAIAGI": None,
               "KXBA": "BA", "KXBOEING": "BA", "KXMETADAP": "META", "KXMETAHEADCOUNT": "META", "KXPALANTIR": "PLTR",
               "KXRIVN": "RIVN", "KXSNOWFLAKE": "SNOW", "KXSOFIMEMBERS": "SOFI", "KXSPOTIFYSUBS": "SPOT", "KXTESLA": "TSLA",
               "KXTESLAPROD": "TSLA", "KXTSLA": "TSLA", "KXUBERTRIPS": "UBER", "KXDASHORDERS": "DASH", "KXRHGOLD": "HOOD",
               "KXMATCHPAYERS": "MTCH", "KXPBR": "PBR"}   # each is still checked against Yahoo earnings dates

_S = requests.Session()
_S.headers.update({"User-Agent": "IUM-BeatOdds-student-project/1.0"})


def _get(url, params=None, tries=3):
    for i in range(tries):
        try:
            r = _S.get(url, params=params, timeout=30)
            if r.status_code == 429:
                time.sleep(2 + 3 * i); continue
            if r.status_code == 404:
                return None
            r.raise_for_status()
            return r.json()
        except Exception:
            if i == tries - 1:
                raise
            time.sleep(1 + i)


def _jl(x):
    return json.loads(x) if isinstance(x, str) else (x or [])


# =============================================================================
# 1. Polymarket earnings markets
# =============================================================================
def fetch_events(closed):
    out, off = [], 0
    while True:
        page = _get(GAMMA + "/events", {"tag_slug": "earnings", "closed": str(closed).lower(), "limit": 100, "offset": off}) or []
        out += page
        if len(page) < 100:
            break
        off += 100
        time.sleep(0.2)
    return out


def _parse_market(ev, m):
    """One Polymarket earnings contract -> ticker, report date, EPS threshold, GAAP or not, outcome. None if unusable."""
    q = (m.get("question") or ev.get("title") or "")
    if "beat" not in q.lower():
        return None
    slug = m.get("slug") or ev.get("slug") or ""
    tk = re.search(r"\(([A-Z][A-Z.\-]{0,6})\)", q) or re.search(r"\(([A-Z][A-Z.\-]{0,6})\)", ev.get("title") or "")
    ticker = tk.group(1) if tk else None
    if not ticker:
        s0 = re.match(r"^([a-z]{1,6})-quarterly-earnings", slug)
        ticker = s0.group(1).upper() if s0 else None
    d = (re.search(r"(\d{2})-(\d{2})-(\d{4})", slug), re.search(r"(\d{4})-(\d{2})-(\d{2})", slug))
    if d[0]:
        rdate = pd.Timestamp(f"{d[0].group(3)}-{d[0].group(1)}-{d[0].group(2)}")
    elif d[1]:
        rdate = pd.Timestamp(d[1].group(0))
    else:
        rdate = pd.Timestamp(str(m.get("endDate") or ev.get("endDate"))[:10]) if (m.get("endDate") or ev.get("endDate")) else None
    th = re.search(r"(neg-?|m)?(\d+)pt(\d+)$", slug)
    if th:
        thr = float(f"{th.group(2)}.{th.group(3)}") * (-1 if th.group(1) else 1)
    else:
        tt = re.search(r"\$\s?(-?\d+(?:\.\d+)?)", q + " " + (ev.get("description") or ""))
        thr = float(tt.group(1)) if tt else None
    outs, prices = [o.lower() for o in _jl(m.get("outcomes"))], _jl(m.get("outcomePrices"))
    toks = _jl(m.get("clobTokenIds"))
    if "yes" not in outs or not toks:
        return None
    iy = outs.index("yes")
    try:
        p_last = float(prices[iy])
    except Exception:
        p_last = None
    closed = bool(m.get("closed"))
    resolved = closed and (str(m.get("umaResolutionStatus", "")).lower() == "resolved" or p_last in (0.0, 1.0)) and p_last in (0.0, 1.0)
    return dict(slug=slug, question=q, ticker=ticker, report_date=rdate, threshold=thr,
                eps_type="non-GAAP" if "nongaap" in slug else ("GAAP" if "gaap" in slug else "not stated"),
                token=toks[iy], volume=float(m.get("volume") or 0), created=str(m.get("startDate") or ev.get("startDate") or "")[:19],
                closed=closed, resolved=resolved, outcome=(int(p_last) if resolved else None), p_now=p_last)


def fetch_markets():
    rows = []
    for closed in (True, False):
        evs = fetch_events(closed)
        print(f"   {'closed' if closed else 'open'} earnings events: {len(evs)}", flush=True)
        for ev in evs:
            for m in ev.get("markets", []) or []:
                x = _parse_market(ev, m)
                if x:
                    rows.append(x)
    df = pd.DataFrame(rows).drop_duplicates("slug")
    return df


def fetch_odds(token, rdate):
    """Hourly YES price over the last ODDS_WINDOW_DAYS before the cut-off on the report day."""
    cut = int(datetime(rdate.year, rdate.month, rdate.day, CUTOFF_HOUR_UTC, tzinfo=timezone.utc).timestamp())
    try:
        js = _get(CLOB + "/prices-history", {"market": token, "startTs": cut - ODDS_WINDOW_DAYS * 86400, "endTs": cut + 3600,
                                             "fidelity": 60}) or {}
        return [(int(p["t"]), float(p["p"])) for p in js.get("history", []) or [] if int(p["t"]) <= cut]
    except Exception:
        return []


# =============================================================================
# 1b. Kalshi KPI markets
# =============================================================================
def _ts(s):
    try:
        return int(pd.Timestamp(s).timestamp())
    except Exception:
        return None


def _f(x):
    try:
        return float(x)
    except Exception:
        return None


def k_series():
    """All Kalshi series tagged as company KPIs."""
    out = {}
    for params in ({"tags": "KPIs"}, {"category": "Financials"}, {"category": "Companies"}):
        try:
            js = _get(KALSHI + "/series", params) or {}
        except Exception:
            continue
        for s in js.get("series") or []:
            tags = [str(t).lower() for t in (s.get("tags") or [])]
            if "kpis" in tags or "kpi" in str(s.get("title", "")).lower():
                out[s["ticker"]] = s.get("title", "")
    return out


def k_markets(series):
    """Settled markets of one series: the historical archive plus the recent ones."""
    rows = []
    for url, extra in ((KALSHI + "/historical/markets", {}), (KALSHI + "/markets", {"status": "settled"})):
        cursor = None
        for _ in range(50):
            params = {"series_ticker": series, "limit": 1000, **extra}
            if cursor:
                params["cursor"] = cursor
            try:
                js = _get(url, params, tries=2) or {}
            except Exception:
                break
            rows += js.get("markets") or []
            cursor = js.get("cursor")
            if not cursor or not js.get("markets"):
                break
    return rows


def _k_parse(series, m):
    """One Kalshi KPI contract -> strike, outcome, actual value. Only 'above X' contracts are kept."""
    res = str(m.get("result") or "").lower()
    if res not in ("yes", "no"):
        return None
    fs, cap = _f(m.get("floor_strike")), m.get("cap_strike")
    st = str(m.get("strike_type") or "").lower()
    sub = str(m.get("yes_sub_title") or m.get("subtitle") or "").lower()
    title = str(m.get("title") or "")
    above = st in ("greater", "greater_or_equal") or (st in ("structured", "custom", "") and
                                                      (sub.startswith("above") or sub.startswith("at least") or " above " in title.lower()))
    cts = _ts(m.get("close_time"))
    if fs is None or cap is not None or not above or not cts:
        return None
    return dict(mticker=m.get("ticker"), series=series, event=m.get("event_ticker"), title=title, strike=fs,
                actual=_f(m.get("expiration_value")), y=1 if res == "yes" else 0,
                volume=_f(m.get("volume_fp")) or _f(m.get("volume")) or 0.0, close_ts=cts)


def k_candidates(series):
    if series in K_OVERRIDES:
        return [K_OVERRIDES[series]] if K_OVERRIDES[series] else []
    s = series[2:] if series.startswith("KX") else series
    c = [s] + ([s[:-1]] if s.endswith("A") and len(s) > 2 else [])
    return [x for x in c if 1 <= len(x) <= 5]


def _k_price(c):
    """Last traded YES price in a candle; mid quote if nothing traded yet. Handles dollar strings and old cent integers."""
    def g(d, *keys):
        for k in keys:
            v = _f((d or {}).get(k))
            if v is not None:
                return v / 100 if v > 1.0001 else v
        return None
    p = g(c.get("price"), "close_dollars", "close", "previous_dollars", "previous")
    if p is None:
        b, a = g(c.get("yes_bid"), "close_dollars", "close"), g(c.get("yes_ask"), "close_dollars", "close")
        p = (a + b) / 2 if a is not None and b is not None and a - b <= 0.2 else None
    return p


def k_odds(row, start=None):
    """Hourly YES price over the last ODDS_WINDOW_DAYS before the cut-off. Tries the archive and the live endpoints."""
    t, cut = row["mticker"], int(row["cut"])
    p = {"start_ts": start or (cut - ODDS_WINDOW_DAYS * 86400), "end_ts": cut + 3600, "period_interval": 60}
    for url, extra, key in ((f"{KALSHI}/historical/markets/{t}/candlesticks", {}, None),
                            (f"{KALSHI}/markets/candlesticks", {"market_tickers": t}, "markets"),
                            (f"{KALSHI}/series/{row['series']}/markets/{t}/candlesticks", {}, None)):
        try:
            js = _get(url, {**p, **extra}, tries=3) or {}
        except Exception:
            continue
        cs = (js.get("markets") or [{}])[0].get("candlesticks", []) if key else js.get("candlesticks", [])
        pts = [(int(c["end_period_ts"]), _k_price(c)) for c in cs or [] if c.get("end_period_ts")]
        pts = [(a, b) for a, b in pts if b is not None and a <= cut]
        if pts:
            return sorted(pts)
    return []


def fetch_kalshi():
    """Settled 'above X' KPI contracts since START_DATE, one per KPI event (the most traded strike)."""
    ser = k_series()
    print(f"   Kalshi KPI series: {len(ser)}", flush=True)
    with ThreadPoolExecutor(THREADS) as ex:
        raw = list(ex.map(k_markets, list(ser)))
    rows = [x for s, ms in zip(ser, raw) for m in ms for x in [_k_parse(s, m)] if x]
    if not rows:
        return pd.DataFrame()
    df = pd.DataFrame(rows).drop_duplicates("mticker")
    df = df[(df["close_ts"] >= _ts(START_DATE)) & (df["volume"] >= K_MIN_VOLUME)]
    df = df.sort_values("volume", ascending=False).drop_duplicates("event")
    print(f"   settled KPI contracts: {len(rows)}, KPI events kept (most traded strike): {len(df)}", flush=True)
    return df.reset_index(drop=True)


def k_resolve(km, est_by):
    """Stock ticker for each series (the candidate whose Yahoo earnings dates match the contract close dates) and the report date."""
    pick = {}
    for s, g in km.groupby("series"):
        best, score = None, 0
        days = [pd.Timestamp(c - 5 * 3600, unit="s").normalize() for c in g["close_ts"]]
        for c in k_candidates(s):
            ds = [x["date"] for x in est_by.get(c, [])]
            n = sum(any(abs((d - e).days) <= 2 for e in ds) for d in days)
            if n > score:
                best, score = c, n
        pick[s] = best
    out = []
    for r in km.to_dict("records"):
        t = pick.get(r["series"])
        if not t:
            continue
        d0 = pd.Timestamp(r["close_ts"] - 5 * 3600, unit="s").normalize()
        near = [x["date"] for x in est_by.get(t, []) if abs((x["date"] - d0).days) <= 2]
        rd = min(near, key=lambda e: abs((e - d0).days)) if near else d0
        cut = min(int(datetime(rd.year, rd.month, rd.day, CUTOFF_HOUR_UTC, tzinfo=timezone.utc).timestamp()), r["close_ts"] - 3600)
        out.append({**r, "ticker": t, "report_date": rd, "cut": cut, "date_matched": bool(near)})
    miss = sorted(s for s, t in pick.items() if not t)
    if miss:
        print(f"   Kalshi series with no stock ticker found (add to K_OVERRIDES if listed): {', '.join(miss)}", flush=True)
    return pd.DataFrame(out)


# =============================================================================
# 2. Yahoo Finance: prices and analyst estimates
# =============================================================================
def yf_symbol(t):
    return t.replace(".", "-")


def fetch_prices(tickers, start):
    import yfinance as yf
    syms = sorted({yf_symbol(t) for t in tickers} | {MARKET})
    data = yf.download(syms, start=str((pd.Timestamp(start) - timedelta(days=15)).date()),
                       auto_adjust=True, progress=False, group_by="column", threads=True)
    px = data["Close"] if isinstance(data.columns, pd.MultiIndex) else data[["Close"]].rename(columns={"Close": syms[0]})
    px.index = pd.to_datetime(px.index).tz_localize(None).normalize()
    return px.dropna(how="all")


def fetch_estimates(ticker):
    import logging
    import yfinance as yf
    logging.getLogger("yfinance").setLevel(logging.CRITICAL)
    try:
        e = yf.Ticker(yf_symbol(ticker)).get_earnings_dates(limit=20)
        if e is None or e.empty:
            return []
        e = e.reset_index()
        dcol = e.columns[0]
        out = []
        for _, r in e.iterrows():
            est, rep = r.get("EPS Estimate"), r.get("Reported EPS")
            if pd.notna(est) and pd.notna(rep):
                out.append(dict(ticker=ticker, date=pd.Timestamp(r[dcol]).tz_localize(None).normalize(),
                                estimate=float(est), reported=float(rep),
                                surprise_pct=float(r["Surprise(%)"]) if pd.notna(r.get("Surprise(%)")) else None))
        return out
    except Exception:
        return []


def fetch_option_move(ticker, rdate):
    """ATM straddle on the first expiry after the report date: the move the options market prices. Today's snapshot only."""
    import logging
    import yfinance as yf
    logging.getLogger("yfinance").setLevel(logging.CRITICAL)
    try:
        t = yf.Ticker(yf_symbol(ticker))
        exps = [e for e in (t.options or []) if pd.Timestamp(e) > pd.Timestamp(rdate)]
        if not exps:
            return None
        spot = float(t.history(period="5d")["Close"].dropna().iloc[-1])
        ch = t.option_chain(exps[0])
        def mid(df, k):
            x = df[df["strike"] == k]
            if x.empty:
                return None
            b, a, l = (float(x.iloc[0][c]) if pd.notna(x.iloc[0][c]) else 0.0 for c in ("bid", "ask", "lastPrice"))
            return (a + b) / 2 if a > 0 and b > 0 else (l or None)
        ks = sorted(set(ch.calls["strike"]) & set(ch.puts["strike"]), key=lambda k: abs(k - spot))
        if not ks:
            return None
        c, p = mid(ch.calls, ks[0]), mid(ch.puts, ks[0])
        if c is None or p is None:
            return None
        return dict(ticker=ticker, expiry=pd.Timestamp(exps[0]), spot=spot, strike=float(ks[0]), call=c, put=p)
    except Exception:
        return None


# =============================================================================
# 2b. SEC EDGAR release times and FINRA short interest
# =============================================================================
import threading
_SEC = requests.Session()
_SEC_LOCK = threading.Lock()
_SEC_LAST = [0.0]


def _sec_get(url):
    """SEC asks for at most 10 requests per second and a User-Agent with a contact email."""
    _SEC.headers.update({"User-Agent": SEC_USER_AGENT, "Accept-Encoding": "gzip, deflate"})
    for i in range(3):
        with _SEC_LOCK:
            wait = 0.12 - (time.time() - _SEC_LAST[0])
            if wait > 0:
                time.sleep(wait)
            _SEC_LAST[0] = time.time()
        try:
            r = _SEC.get(url, timeout=30)
            if r.status_code in (403, 429):
                time.sleep(2 + 2 * i); continue
            r.raise_for_status()
            return r.json()
        except Exception:
            time.sleep(1 + i)
    return None


def sec_cik_map():
    js = _sec_get("https://www.sec.gov/files/company_tickers.json") or {}
    return {str(v["ticker"]).upper().replace("-", "."): int(v["cik_str"]) for v in js.values()}


def _edgar_rows(block, ticker):
    keys = ("form", "filingDate", "acceptanceDateTime", "items")
    if not all(k in block for k in keys):
        return []
    out = []
    for f, d, a, it in zip(*(block[k] for k in keys)):
        if str(f).startswith("8-K") and "2.02" in str(it or ""):
            out.append(dict(ticker=ticker, filed=d, acceptance=a))
    return out


def fetch_edgar(ticker, cik, dmin):
    """Earnings press-release filings (8-K, item 2.02) of one company since dmin."""
    js = _sec_get(f"https://data.sec.gov/submissions/CIK{cik:010d}.json")
    if not js:
        return []
    rec = (js.get("filings") or {}).get("recent") or {}
    rows = _edgar_rows(rec, ticker)
    dates = rec.get("filingDate") or []
    if dates and min(dates) > dmin:
        for f in ((js.get("filings") or {}).get("files") or [])[:8]:
            if str(f.get("filingTo", "")) >= dmin:
                more = _sec_get("https://data.sec.gov/submissions/" + f["name"])
                if more:
                    rows += _edgar_rows(more, ticker)
    return [r for r in rows if r["filed"] >= dmin]


def edgar_time(acc):
    """EDGAR labels times with Z, but some are Eastern wall time. Earnings releases come out before about 09:30
    or after 16:00 New York time, so a time that would fall at night or in the afternoon is read as Eastern instead."""
    try:
        t = pd.Timestamp(str(acc).replace("Z", "").replace(".000", ""))
        u = t.tz_localize("UTC"); e = u.tz_convert("America/New_York"); note = "read as UTC"
        h = e.hour + e.minute / 60
        if h < 5.5 or 11.5 <= h < 15.75:
            e = t.tz_localize("America/New_York", ambiguous="NaT", nonexistent="shift_forward")
            u = e.tz_convert("UTC"); note = "read as New York time"
        h = e.hour + e.minute / 60
        ses = "before open" if h < 9.5 else ("after close" if h >= 16 else "during market")
        return u.tz_localize(None), e.tz_localize(None), ses, note
    except Exception:
        return None, None, None, None


def fetch_all_edgar(tickers, dmin):
    try:
        cmap = sec_cik_map()
    except Exception:
        cmap = {}
    pairs = [(t, cmap[t.upper()]) for t in tickers if t.upper() in cmap]
    with ThreadPoolExecutor(4) as ex:
        rows = [r for rs in ex.map(lambda x: fetch_edgar(x[0], x[1], dmin), pairs) for r in rs]
    df = pd.DataFrame(rows)
    if len(df):
        tt = [edgar_time(a) for a in df["acceptance"]]
        df["release_utc"], df["release_et"], df["session"], df["read_as"] = zip(*tt)
        df = df.dropna(subset=["release_utc"]).sort_values(["ticker", "release_utc"]).reset_index(drop=True)
    print(f"   EDGAR: {len(pairs)} of {len(tickers)} tickers found, {len(df)} earnings filings", flush=True)
    return df


def _finra_post(body):
    url = "https://api.finra.org/data/group/otcMarket/name/consolidatedShortInterest"
    for i in range(3):
        try:
            r = _S.post(url, json=body, headers={"Accept": "application/json", "Content-Type": "application/json"}, timeout=60)
            if r.status_code == 204:
                return []
            if r.status_code == 429:
                time.sleep(3 + 3 * i); continue
            r.raise_for_status()
            return r.json()
        except Exception:
            time.sleep(1 + i)
    return None


def fetch_short(tickers, dmin, dmax):
    """FINRA consolidated short interest for our tickers, all settlement dates in the window."""
    rows, syms = [], sorted({t.upper() for t in tickers})
    dr = [{"fieldName": "settlementDate", "startDate": dmin, "endDate": dmax}]
    for k in range(0, len(syms), 100):
        chunk, off = syms[k:k + 100], 0
        while True:
            js = _finra_post({"limit": 5000, "offset": off, "dateRangeFilters": dr,
                              "domainFilters": [{"fieldName": "symbolCode", "values": chunk}]})
            if js is None:          # filter not accepted: one request per symbol instead
                for sy in chunk:
                    js2 = _finra_post({"limit": 500, "dateRangeFilters": dr,
                                       "compareFilters": [{"fieldName": "symbolCode", "compareType": "EQUAL", "fieldValue": sy}]})
                    rows += js2 or []
                break
            rows += js
            if len(js) < 5000:
                break
            off += 5000
    df = pd.DataFrame(rows)
    if len(df):
        df = df.rename(columns={"symbolCode": "symbol", "settlementDate": "settle", "currentShortPositionQuantity": "short",
                                "averageDailyVolumeQuantity": "adv", "daysToCoverQuantity": "dtc"})
        keep = [c for c in ("symbol", "settle", "short", "adv", "dtc") if c in df.columns]
        df = df[keep].copy()
        df["settle"] = pd.to_datetime(df["settle"])
        for c in ("short", "adv", "dtc"):
            df[c] = pd.to_numeric(df[c], errors="coerce")
        df = df[df["symbol"].isin(syms)].drop_duplicates(["symbol", "settle"]).sort_values(["symbol", "settle"]).reset_index(drop=True)
    print(f"   FINRA short interest rows: {len(df)}", flush=True)
    return df


# =============================================================================
# 2c. Fed meetings: Kalshi, Polymarket and FRED
# =============================================================================
FED_LABELS = ("hold", "cut 25bp", "cut more than 25bp", "hike 25bp", "hike more than 25bp")


def _fed_outcome(text):
    """Map a contract label to one of five outcomes."""
    t = str(text).lower().replace("basis points", "bps").replace(" bps", "bps").replace(" bp", "bps")
    if "maintain" in t or "no change" in t or "unchanged" in t or re.search(r"(^|[^0-9])0bps", t):
        return "hold"
    big = (">25" in t) or ("50" in t) or ("more than 25" in t) or ("75" in t)
    if "cut" in t or "decrease" in t or "lower" in t:
        return "cut more than 25bp" if big else "cut 25bp"
    if "hike" in t or "increase" in t or "raise" in t:
        return "hike more than 25bp" if big else "hike 25bp"
    return None


def fetch_fed_kalshi():
    rows = []
    for ser in FED_SERIES:
        for url, extra in ((KALSHI + "/historical/markets", {}), (KALSHI + "/markets", {"status": "settled"}), (KALSHI + "/markets", {"status": "open"})):
            cursor = None
            for _ in range(20):
                params = {"series_ticker": ser, "limit": 1000, **extra}
                if cursor:
                    params["cursor"] = cursor
                try:
                    js = _get(url, params, tries=2) or {}
                except Exception:
                    break
                for m in js.get("markets") or []:
                    lab = _fed_outcome(m.get("yes_sub_title") or m.get("subtitle") or m.get("title"))
                    cts = _ts(m.get("close_time"))
                    if not lab or not cts:
                        continue
                    res = str(m.get("result") or "").lower()
                    rows.append(dict(venue="Kalshi", series=ser, market=m.get("ticker"), meeting=m.get("event_ticker"), outcome=lab,
                                     y=1 if res == "yes" else (0 if res == "no" else None), close_ts=cts,
                                     volume=_f(m.get("volume_fp")) or _f(m.get("volume")) or 0.0,
                                     p_now=_f(m.get("last_price_dollars")) if m.get("last_price_dollars") is not None else None,
                                     open_=str(m.get("status", "")).lower() in ("open", "active", "initialized")))
                cursor = js.get("cursor")
                if not cursor or not js.get("markets"):
                    break
    df = pd.DataFrame(rows)
    if len(df):
        df = df.drop_duplicates("market")
        df["date"] = pd.to_datetime(df["close_ts"], unit="s").dt.normalize()
        df = df[df["date"] >= pd.Timestamp(FED_START)]
    return df


def fetch_fed_poly():
    rows, seen = [], set()
    for closed in (True, False):
        for tag in FED_TAGS:
            off = 0
            for _ in range(10):
                try:
                    page = _get(GAMMA + "/events", {"tag_slug": tag, "closed": str(closed).lower(), "limit": 100, "offset": off}) or []
                except Exception:
                    break
                for ev in page:
                    title = str(ev.get("title") or "")
                    if "fed" not in title.lower() or "decision" not in title.lower() or ev.get("id") in seen:
                        continue
                    seen.add(ev.get("id"))
                    for m in ev.get("markets") or []:
                        qq = f"{m.get('question') or ''} {m.get('slug') or ''}".lower()
                        if "meeting" not in qq or "dissent" in qq:
                            continue
                        lab = _fed_outcome(m.get("groupItemTitle") or m.get("question"))
                        outs, prices, toks = [o.lower() for o in _jl(m.get("outcomes"))], _jl(m.get("outcomePrices")), _jl(m.get("clobTokenIds"))
                        if not lab or "yes" not in outs or not toks:
                            continue
                        iy = outs.index("yes")
                        try:
                            pl = float(prices[iy])
                        except Exception:
                            pl = None
                        end = m.get("endDate") or ev.get("endDate")
                        cts = _ts(end)
                        if not cts:
                            continue
                        res = int(pl) if (m.get("closed") and pl in (0.0, 1.0)) else None
                        rows.append(dict(venue="Polymarket", series="", market=m.get("slug"), meeting=f"{ev.get('slug')}|{str(end)[:10]}", outcome=lab, y=res,
                                         close_ts=cts, volume=float(m.get("volume") or 0), p_now=pl, open_=not m.get("closed"), token=toks[iy]))
                if len(page) < 100:
                    break
                off += 100
    df = pd.DataFrame(rows)
    if len(df):
        df = df.drop_duplicates("market")
        df["date"] = pd.to_datetime(df["close_ts"], unit="s").dt.normalize()
        df = df[df["date"] >= pd.Timestamp(FED_START)]
        # several markets with the same outcome in one meeting: keep the most traded
        df = df.sort_values("volume", ascending=False).drop_duplicates(["meeting", "outcome"])
    return df


def fed_odds(row):
    """Hourly price of one Fed outcome in three short windows: around 30, 7 and 1 days before the decision."""
    end = int(row["close_ts"]); pts = []
    for back in (31, 8, 2):
        a, b = end - back * 86400, end - (back - 2) * 86400 if back > 2 else end
        if row["venue"] == "Polymarket":
            try:
                js = _get(CLOB + "/prices-history", {"market": row["token"], "startTs": a, "endTs": b, "fidelity": 60}) or {}
                pts += [(int(x["t"]), float(x["p"])) for x in js.get("history", []) or []]
            except Exception:
                pass
        else:
            pts += k_odds({"mticker": row["market"], "cut": b - 60, "series": row["series"]}, start=a)
    return sorted(set(pts))


def fetch_fred(series, start):
    """FRED public CSV download (no key). Some networks block scripts, so a browser-style header is sent."""
    import io
    try:
        r = requests.get(FRED_CSV, params={"id": series, "cosd": start}, timeout=30,
                         headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)", "Accept": "text/csv,*/*"})
        r.raise_for_status()
        if not r.text.lstrip().lower().startswith(("observation_date", "date")):
            raise ValueError("not a CSV")
        df = pd.read_csv(io.StringIO(r.text))
        df.columns = ["date", "value"]
        df["date"] = pd.to_datetime(df["date"]); df["value"] = pd.to_numeric(df["value"], errors="coerce")
        return df.dropna()
    except Exception:
        return pd.DataFrame(columns=["date", "value"])


def fetch_yield_yahoo(start):
    """Fallback: 5-year Treasury yield (^FVX, in %) from Yahoo Finance."""
    import yfinance as yf
    try:
        x = yf.download("^FVX", start=start, progress=False, auto_adjust=False)["Close"]
        x = x.iloc[:, 0] if hasattr(x, "columns") else x
        df = x.dropna().reset_index(); df.columns = ["date", "value"]
        df["date"] = pd.to_datetime(df["date"]).dt.tz_localize(None)
        return df
    except Exception:
        return pd.DataFrame(columns=["date", "value"])


def fetch_fed():
    k = fetch_fed_kalshi(); pm = fetch_fed_poly()
    fed = pd.concat([x for x in (k, pm) if len(x)], ignore_index=True) if (len(k) or len(pm)) else pd.DataFrame()
    print(f"   Fed outcome contracts: Kalshi {len(k)}, Polymarket {len(pm)}", flush=True)
    odds = []
    if len(fed):
        done = fed[fed["y"].notna()].copy()
        with ThreadPoolExecutor(THREADS) as ex:
            pts = list(ex.map(fed_odds, done.to_dict("records")))
        odds = list(zip(done["market"], pts))
    y2 = fetch_fred("DGS2", str((pd.Timestamp(FED_START) - timedelta(days=60)).date())); yname = "2-year Treasury yield (FRED DGS2)"
    if not len(y2):
        y2 = fetch_yield_yahoo(str((pd.Timestamp(FED_START) - timedelta(days=60)).date())); yname = "5-year Treasury yield (Yahoo ^FVX, FRED unavailable)"
    tgt = fetch_fred("DFEDTARU", str((pd.Timestamp(FED_START) - timedelta(days=60)).date()))
    print(f"   FRED: 2-year yield rows {len(y2)}, Fed target rows {len(tgt)}", flush=True)
    return dict(markets=fed, odds=odds, y2=y2, target=tgt, yname=yname)


# =============================================================================
# 3. Workbook
# =============================================================================
def fix_locale(wb):
    """Write numbers in text criteria as real numbers (">="&0.8, not ">=0.8"), so COUNTIFS and
    AVERAGEIFS also work in Excel set to a comma decimal separator (Italian, French, Dutch, German)."""
    pat = re.compile(r'"(>=|<=|<>|>|<|=)(-?\d*\.\d+)"')
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("=") and '"' in v:
                    c.value = pat.sub(r'"\1"&\2', v)
                elif isinstance(getattr(v, "text", None), str):
                    v.text = pat.sub(r'"\1"&\2', v.text)


def build(markets, odds, prices, est, live, download_utc, km=None, kodds=None, opts=None, hist_px=None, edgar=None, short=None, fed=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as L
    from openpyxl.chart import ScatterChart, Reference, Series, BarChart
    BOLD = Font(name="Arial", bold=True); NOTE = Font(name="Arial", italic=True, color="666666")
    TITLE = Font(name="Arial", bold=True, size=14, color="1F3A5F"); SEC = Font(name="Arial", bold=True, size=11, color="1F3A5F")
    HEAD = PatternFill("solid", fgColor="DDE4EE"); YELLOW = PatternFill("solid", fgColor="FFFF00"); BLUE = Font(name="Arial", color="0000FF")

    def head(ws, r, vals, c0=1):
        for j, v in enumerate(vals):
            c = ws.cell(row=r, column=c0 + j, value=v); c.font = BOLD; c.fill = HEAD
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 45

    wb = Workbook(); wb.remove(wb.active)
    RES, EVT, LIV, ASM = (wb.create_sheet(n) for n in ("Results", "Events", "Live", "Assumptions"))
    RM, RO, RP, RE, RL = (wb.create_sheet(n) for n in ("Raw_Markets", "Raw_Odds", "Raw_Prices", "Raw_Estimates", "Raw_Live"))

    # ---------- raw sheets ----------
    RM["A1"] = f"Polymarket earnings contracts as downloaded: {GAMMA}/events?tag_slug=earnings. Downloaded {download_utc} UTC."; RM["A1"].font = BOLD
    cols = ["slug", "question", "ticker", "report_date", "threshold", "eps_type", "volume", "created", "outcome", "token"]
    head(RM, 4, ["Contract slug", "Question", "Ticker", "Report date", "EPS threshold (consensus locked at listing)", "EPS type",
                 "Volume (USD)", "Listed (UTC)", "Outcome (1 = beat, 0 = miss)", "CLOB token (YES)"])
    for i, r in markets.reset_index(drop=True).iterrows():
        for j, c in enumerate(cols):
            v = r[c]
            if c == "report_date":
                v = r[c].to_pydatetime()
            RM.cell(row=5 + i, column=1 + j, value=v)
        RM.cell(row=5 + i, column=4).number_format = "yyyy-mm-dd"
    RO["A1"] = f"Hourly YES prices as downloaded: {CLOB}/prices-history (fidelity 60), last {ODDS_WINDOW_DAYS} days before the cut-off. Sorted by contract and time."
    RO["A1"].font = BOLD
    head(RO, 4, ["Contract row (Raw_Markets)", "Time (unix, UTC)", "Time (UTC)", "YES price"])
    r = 5
    span = {}
    for i, (slug, pts) in enumerate(odds):
        a = r
        for t, p in pts:
            RO.cell(row=r, column=1, value=5 + i); RO.cell(row=r, column=2, value=t)
            RO.cell(row=r, column=3, value=datetime.fromtimestamp(t, timezone.utc).replace(tzinfo=None)).number_format = "yyyy-mm-dd hh:mm"
            RO.cell(row=r, column=4, value=p)
            r += 1
        span[slug] = (a, r - 1) if r > a else None
    RP["A1"] = f"Daily adjusted close as downloaded: Yahoo Finance via yfinance. Downloaded {download_utc} UTC."; RP["A1"].font = BOLD
    head(RP, 4, ["Date"] + list(prices.columns))
    for i, (d, row) in enumerate(prices.iterrows()):
        RP.cell(row=5 + i, column=1, value=d.to_pydatetime()).number_format = "yyyy-mm-dd"
        for j, c in enumerate(prices.columns):
            if pd.notna(row[c]):
                RP.cell(row=5 + i, column=2 + j, value=float(row[c]))
    PE = 4 + len(prices); PL = L(1 + len(prices.columns))
    RE["A1"] = f"Analyst EPS estimates and reported EPS as downloaded: Yahoo Finance via yfinance (get_earnings_dates). Downloaded {download_utc} UTC."
    RE["A1"].font = BOLD
    head(RE, 4, ["Ticker", "Report date", "EPS estimate", "Reported EPS", "Surprise (%) as published", "Beat (1/0)"])
    for i, x in enumerate(est):
        rr = 5 + i
        RE.cell(row=rr, column=1, value=x["ticker"]); RE.cell(row=rr, column=2, value=x["date"].to_pydatetime()).number_format = "yyyy-mm-dd"
        RE.cell(row=rr, column=3, value=x["estimate"]); RE.cell(row=rr, column=4, value=x["reported"])
        if x["surprise_pct"] is not None:
            RE.cell(row=rr, column=5, value=x["surprise_pct"])
        RE.cell(row=rr, column=6, value=f"=IF(D{rr}>C{rr},1,0)")
    EE = max(5, 4 + len(est))
    RL["A1"] = f"Open earnings contracts as downloaded: {GAMMA}/events (closed=false). Current YES price = outcomePrices. Downloaded {download_utc} UTC."
    RL["A1"].font = BOLD
    head(RL, 4, ["Contract slug", "Question", "Ticker", "Report date", "EPS threshold", "EPS type", "Volume (USD)", "YES price now"])
    for i, x in enumerate(live):
        for j, c in enumerate(["slug", "question", "ticker", "report_date", "threshold", "eps_type", "volume", "p_now"]):
            v = x[c].to_pydatetime() if c == "report_date" else x[c]
            RL.cell(row=5 + i, column=1 + j, value=v)
        RL.cell(row=5 + i, column=4).number_format = "yyyy-mm-dd"
    for ws in (RM, RO, RP, RE, RL):
        ws.freeze_panes = "A5"

    # ---------- Events: one row per usable resolved contract ----------
    EVT["A1"] = "Events: one row per resolved earnings contract with odds and prices (all cells are formulas linking to the Raw_ sheets)"
    EVT["A1"].font = TITLE
    EVT["A2"] = (f"p = YES price at {CUTOFF_HOUR_UTC}:00 UTC on the report day (before the release, whether before the open or after the close). "
                 "Reaction = log return from the close before the report day to the close after it, which covers both release times. "
                 "Abnormal = stock minus " + MARKET + ". Odds rows = first and last row of this contract on Raw_Odds (set by the script).")
    EVT["A2"].alignment = Alignment(wrap_text=True); EVT.merge_cells("A2:Y2"); EVT.row_dimensions[2].height = 45
    hdr = ["Raw row", "Ticker", "Report date", "EPS threshold", "Outcome y (1 = beat)", "Odds first row", "Odds last row",
           "Cut-off (unix)", "p: P(beat) at cut-off", f"p {LEAD_DAYS} days earlier", "Volume (USD)",
           "Running base rate (earlier events)", "Company past beat rate (Yahoo history)", "Analyst surprise (%), capped at +/-100",
           "Brier: Polymarket", "Brier: base rate", "Brier: company history", "All three forecasts available",
           "Stock return (log)", "Market return (log)", "Abnormal return", "Polymarket surprise y - p", "Uncertainty p(1-p)",
           "|Abnormal return|", "Correct side (p>0.5 matches y)", f"Brier: Polymarket {LEAD_DAYS} days earlier"]
    head(EVT, 4, hdr)
    good = []
    for i, x in enumerate(markets.reset_index(drop=True).itertuples()):
        sp = span.get(x.slug)
        if x.resolved and sp and x.ticker and yf_symbol(x.ticker) in prices.columns:
            good.append((5 + i, sp))
    r = 5
    for raw, (a, b) in good:
        f = {}
        f["A"] = raw
        f["B"] = f"=Raw_Markets!C{raw}"; f["C"] = f"=Raw_Markets!D{raw}"; f["D"] = f"=Raw_Markets!E{raw}"; f["E"] = f"=Raw_Markets!I{raw}"
        f["F"] = a; f["G"] = b
        f["H"] = f"=(C{r}-DATE(1970,1,1))*86400+{CUTOFF_HOUR_UTC}*3600"
        f["I"] = f'=IFERROR(INDEX(Raw_Odds!$D${a}:$D${b},MATCH(H{r},Raw_Odds!$B${a}:$B${b},1)),"")'
        f["J"] = f'=IFERROR(INDEX(Raw_Odds!$D${a}:$D${b},MATCH(H{r}-{LEAD_DAYS}*86400,Raw_Odds!$B${a}:$B${b},1)),"")'
        f["K"] = f"=Raw_Markets!G{raw}"
        f["L"] = f'=IF(COUNTIF($C$5:$C$__E__,"<"&C{r})>={MIN_PRIOR_EVENTS},AVERAGEIFS($E$5:$E$__E__,$C$5:$C$__E__,"<"&C{r}),"")'
        f["M"] = (f'=IF(COUNTIFS(Raw_Estimates!$A$5:$A${EE},B{r},Raw_Estimates!$B$5:$B${EE},"<"&C{r}-3)>={MIN_HISTORY},'
                  f'AVERAGEIFS(Raw_Estimates!$F$5:$F${EE},Raw_Estimates!$A$5:$A${EE},B{r},Raw_Estimates!$B$5:$B${EE},"<"&C{r}-3),"")')
        f["N"] = (f'=IFERROR(MAX(-100,MIN(100,AVERAGEIFS(Raw_Estimates!$E$5:$E${EE},Raw_Estimates!$A$5:$A${EE},B{r},'
                  f'Raw_Estimates!$B$5:$B${EE},">="&C{r}-3,Raw_Estimates!$B$5:$B${EE},"<="&C{r}+3))),"")')
        f["O"] = f'=IF(ISNUMBER(I{r}),(I{r}-E{r})^2,"")'
        f["P"] = f'=IF(ISNUMBER(L{r}),(L{r}-E{r})^2,"")'
        f["Q"] = f'=IF(ISNUMBER(M{r}),(M{r}-E{r})^2,"")'
        f["R"] = f"=IF(AND(ISNUMBER(O{r}),ISNUMBER(P{r}),ISNUMBER(Q{r})),1,0)"
        sym = yf_symbol(markets.iloc[raw - 5]["ticker"])
        cS = L(2 + list(prices.columns).index(sym)); cM = L(2 + list(prices.columns).index(MARKET))
        pre = f"MATCH(C{r}-1,Raw_Prices!$A$5:$A${PE},1)"; post = f"MATCH(C{r},Raw_Prices!$A$5:$A${PE},1)+1"
        f["S"] = f'=IFERROR(LN(INDEX(Raw_Prices!${cS}$5:${cS}${PE},{post})/INDEX(Raw_Prices!${cS}$5:${cS}${PE},{pre})),"")'
        f["T"] = f'=IFERROR(LN(INDEX(Raw_Prices!${cM}$5:${cM}${PE},{post})/INDEX(Raw_Prices!${cM}$5:${cM}${PE},{pre})),"")'
        f["U"] = f'=IF(AND(ISNUMBER(S{r}),ISNUMBER(T{r})),S{r}-T{r},"")'
        f["V"] = f'=IF(ISNUMBER(I{r}),E{r}-I{r},"")'
        f["W"] = f'=IF(ISNUMBER(I{r}),I{r}*(1-I{r}),"")'
        f["X"] = f'=IF(ISNUMBER(U{r}),ABS(U{r}),"")'
        f["Y"] = f'=IF(ISNUMBER(I{r}),IF((I{r}>0.5)=(E{r}=1),1,0),"")'
        f["Z"] = f'=IF(ISNUMBER(J{r}),(J{r}-E{r})^2,"")'
        for k, v in f.items():
            EVT[f"{k}{r}"] = v
        r += 1
    E_ = r - 1
    for rr in range(5, r):
        for k in ("L",):
            EVT[f"{k}{rr}"] = EVT[f"{k}{rr}"].value.replace("__E__", str(E_))
        EVT[f"C{rr}"].number_format = "yyyy-mm-dd"
        for k in "IJLMOPQZ":
            EVT[f"{k}{rr}"].number_format = "0.000"
        for k in "STUVX":
            EVT[f"{k}{rr}"].number_format = "+0.00%;-0.00%" if k in "STUX" else "+0.00;-0.00"
        EVT[f"W{rr}"].number_format = "0.000"; EVT[f"K{rr}"].number_format = "#,##0"
    EVT.freeze_panes = "C5"
    R = lambda c: f"Events!${c}$5:${c}${E_}"

    # ---------- Results ----------
    RES["A1"] = "BeatOdds: do Polymarket earnings odds add information beyond consensus and history?"; RES["A1"].font = TITLE
    RES["A2"] = (f"{len(good)} resolved earnings contracts since {START_DATE}. Brier score = mean of (forecast - outcome)^2, lower is better. "
                 "Skill = 1 - Brier(Polymarket) / Brier(benchmark): above 0 means Polymarket beats the benchmark. "
                 "Benchmarks: the running share of beats among earlier events, and the company's own beat rate over its past reports.")
    RES["A2"].alignment = Alignment(wrap_text=True); RES.merge_cells("A2:H2"); RES.row_dimensions[2].height = 45
    rr = 4
    RES.cell(row=rr, column=1, value="A. Accuracy").font = SEC; rr += 1
    rowsA = [("Events", f"=COUNT({R('E')})", "0"),
             ("Share that beat (y = 1)", f"=AVERAGE({R('E')})", "0.0%"),
             ("Average Polymarket p", f"=AVERAGE({R('I')})", "0.0%"),
             ("Correct side (p above 0.5 and beat, or below and miss)", f"=AVERAGE({R('Y')})", "0.0%"),
             ("Always guessing 'beat' would be correct", f"=AVERAGE({R('E')})", "0.0%"),
             ("Events with all three forecasts", f"=SUM({R('R')})", "0"),
             ("Brier, Polymarket (same events)", f"=SUMPRODUCT({R('R')},{R('O')})/SUM({R('R')})", "0.000"),
             ("Brier, running base rate", f"=SUMPRODUCT({R('R')},{R('P')})/SUM({R('R')})", "0.000"),
             ("Brier, company past beat rate", f"=SUMPRODUCT({R('R')},{R('Q')})/SUM({R('R')})", "0.000"),
             ("Skill vs base rate", f"=1-B11/B12", "0.0%"),
             ("Skill vs company history", f"=1-B11/B13", "0.0%"),
             (f"Brier, Polymarket {LEAD_DAYS} days earlier (where available)", f'=IFERROR(AVERAGE({R("Z")}),"n/a")', "0.000")]
    for lab, fm, nf in rowsA:
        RES.cell(row=rr, column=1, value=lab); RES.cell(row=rr, column=2, value=fm).number_format = nf; rr += 1
    rr += 1
    RES.cell(row=rr, column=1, value="Calibration: when Polymarket says p, how often do companies beat?").font = SEC; rr += 1
    head(RES, rr, ["p from", "p to", "Events", "Average p", "Share that beat"]); rr += 1
    cal0 = rr
    for lo, hi in [(0, .2), (.2, .4), (.4, .6), (.6, .8), (.8, .9), (.9, 1.0001)]:
        RES.cell(row=rr, column=1, value=lo); RES.cell(row=rr, column=2, value=hi)
        RES.cell(row=rr, column=3, value=f'=COUNTIFS({R("I")},">="&A{rr},{R("I")},"<"&B{rr})')
        RES.cell(row=rr, column=4, value=f'=IFERROR(AVERAGEIFS({R("I")},{R("I")},">="&A{rr},{R("I")},"<"&B{rr}),"")').number_format = "0.0%"
        RES.cell(row=rr, column=5, value=f'=IFERROR(AVERAGEIFS({R("E")},{R("I")},">="&A{rr},{R("I")},"<"&B{rr}),"")').number_format = "0.0%"
        rr += 1
    cal1 = rr - 1
    ch = ScatterChart(); ch.title = "Calibration: average p vs share that beat"; ch.style = 13; ch.height = 7; ch.width = 12
    ch.x_axis.title = "Average Polymarket p"; ch.y_axis.title = "Share that beat"; ch.x_axis.scaling.min = 0; ch.x_axis.scaling.max = 1
    ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 1; ch.legend = None
    s = Series(Reference(RES, min_col=5, min_row=cal0, max_row=cal1), Reference(RES, min_col=4, min_row=cal0, max_row=cal1))
    s.marker.symbol = "circle"; s.graphicalProperties.line.noFill = True; ch.series.append(s)
    ch.x_axis.delete = False; ch.y_axis.delete = False
    RES.add_chart(ch, "H4")
    rr += 1
    RES.cell(row=rr, column=1, value="B. The stock's reaction: does a surprise relative to Polymarket explain it?").font = SEC; rr += 1
    head(RES, rr, ["Regression of the abnormal return on", "Slope", "t", "R squared", "Events"]); rr += 1
    tt = lambda y, x: f'IFERROR(SLOPE({y},{x})/(STEYX({y},{x})/SQRT(DEVSQ({x}))),"n/a")'
    regs = [("Beat dummy y (1 = beat)", R("E")), ("Polymarket surprise y - p", R("V")), ("Analyst surprise (%)", R("N"))]
    b0 = rr
    for lab, x in regs:
        RES.cell(row=rr, column=1, value=lab)
        RES.cell(row=rr, column=2, value=f'=IFERROR(SLOPE({R("U")},{x}),"n/a")').number_format = "0.0000"
        RES.cell(row=rr, column=3, value="=" + tt(R("U"), x)).number_format = "0.0"
        RES.cell(row=rr, column=4, value=f'=IFERROR(RSQ({R("U")},{x}),"n/a")').number_format = "0.000"
        RES.cell(row=rr, column=5, value=f"=COUNT({x})").number_format = "0"
        rr += 1
    RES.cell(row=rr, column=1, value=("Reading: compare the R squared. If the Polymarket surprise explains more of the reaction than the beat "
                                      "dummy, the odds measure what the market expected better than the consensus does.")).font = NOTE
    rr += 2
    RES.cell(row=rr, column=1, value="C. Is a priced-in beat worth less? Average abnormal return").font = SEC; rr += 1
    head(RES, rr, ["", "Polymarket p at least 0.8", "Polymarket p below 0.8", "Difference"]); rr += 1
    for lab, yv in (("Company beat (y = 1)", 1), ("Company missed (y = 0)", 0)):
        RES.cell(row=rr, column=1, value=lab)
        RES.cell(row=rr, column=2, value=f'=IFERROR(AVERAGEIFS({R("U")},{R("E")},{yv},{R("I")},">=0.8"),"n/a")').number_format = "+0.00%;-0.00%"
        RES.cell(row=rr, column=3, value=f'=IFERROR(AVERAGEIFS({R("U")},{R("E")},{yv},{R("I")},"<0.8"),"n/a")').number_format = "+0.00%;-0.00%"
        RES.cell(row=rr, column=4, value=f'=IFERROR(B{rr}-C{rr},"n/a")').number_format = "+0.00%;-0.00%"
        rr += 1
    RES.cell(row=rr, column=1, value="Count of events in each cell:").font = NOTE
    for j, (yv, op) in enumerate(((1, ">=0.8"), (1, "<0.8"), (0, ">=0.8"), (0, "<0.8"))):
        RES.cell(row=rr, column=2 + j, value=f'=COUNTIFS({R("E")},{yv},{R("I")},"{op}")')
    rr += 2
    RES.cell(row=rr, column=1, value="D. Does uncertain odds mean a bigger move? |Abnormal return| on p(1-p)").font = SEC; rr += 1
    head(RES, rr, ["", "Intercept", "Slope", "t", "R squared"]); rr += 1
    RES.cell(row=rr, column=1, value="|Abnormal return| = a + b x p(1-p)")
    RES.cell(row=rr, column=2, value=f'=IFERROR(INTERCEPT({R("X")},{R("W")}),"n/a")').number_format = "0.0000"
    RES.cell(row=rr, column=3, value=f'=IFERROR(SLOPE({R("X")},{R("W")}),"n/a")').number_format = "0.0000"
    RES.cell(row=rr, column=4, value="=" + tt(R("X"), R("W"))).number_format = "0.0"
    RES.cell(row=rr, column=5, value=f'=IFERROR(RSQ({R("X")},{R("W")}),"n/a")').number_format = "0.000"
    d_row = rr
    rr += 2
    RES.cell(row=rr, column=1, value=("All results are associations over one sample. Odds and prices react to the same information, so "
                                      "the odds are a measure of expectations, not a cause of the stock's move.")).font = NOTE
    RES.column_dimensions["A"].width = 58
    for c in "BCDE":
        RES.column_dimensions[c].width = 16

    # ---------- Live: upcoming reports ----------
    LIV["A1"] = "Live: upcoming earnings with an open Polymarket contract"; LIV["A1"].font = TITLE
    LIV["A2"] = ("p now = current YES price. Company rate = its beat rate over past reports. Expected |move| = intercept + slope x p(1-p) "
                 "from Results part D. Gap = p now minus the company's own history: large gaps mark reports where the crowd disagrees with the past.")
    LIV["A2"].alignment = Alignment(wrap_text=True); LIV.merge_cells("A2:J2"); LIV.row_dimensions[2].height = 40
    head(LIV, 4, ["Ticker", "Report date", "EPS threshold", "EPS type", "Volume (USD)", "p now: P(beat)", "Company past beat rate",
                  "Gap: p now - company rate", "Expected |abnormal move|", "Reading"])
    for i in range(len(live)):
        r, s_ = 5 + i, 5 + i
        LIV[f"A{r}"] = f"=Raw_Live!C{s_}"; LIV[f"B{r}"] = f"=Raw_Live!D{s_}"; LIV[f"B{r}"].number_format = "yyyy-mm-dd"
        LIV[f"C{r}"] = f"=Raw_Live!E{s_}"; LIV[f"D{r}"] = f"=Raw_Live!F{s_}"; LIV[f"E{r}"] = f"=Raw_Live!G{s_}"; LIV[f"E{r}"].number_format = "#,##0"
        LIV[f"F{r}"] = f"=Raw_Live!H{s_}"; LIV[f"F{r}"].number_format = "0%"
        LIV[f"G{r}"] = (f'=IF(COUNTIFS(Raw_Estimates!$A$5:$A${EE},A{r})>={MIN_HISTORY},'
                        f'AVERAGEIFS(Raw_Estimates!$F$5:$F${EE},Raw_Estimates!$A$5:$A${EE},A{r}),"n/a")')
        LIV[f"G{r}"].number_format = "0%"
        LIV[f"H{r}"] = f'=IF(ISNUMBER(G{r}),F{r}-G{r},"n/a")'; LIV[f"H{r}"].number_format = "+0%;-0%"
        LIV[f"I{r}"] = f"=IFERROR(Results!$B${d_row}+Results!$C${d_row}*F{r}*(1-F{r}),\"n/a\")"; LIV[f"I{r}"].number_format = "0.0%"
        LIV[f"J{r}"] = (f'=IF(NOT(ISNUMBER(H{r})),"no company history",IF(H{r}<=-0.2,"crowd more doubtful than history",'
                        f'IF(H{r}>=0.2,"crowd more confident than history","in line with history")))')
    for c, w in zip("ABCDEFGHIJ", (9, 12, 10, 11, 12, 11, 12, 12, 13, 34)):
        LIV.column_dimensions[c].width = w

    # ---------- Assumptions ----------
    ASM["A1"] = "Sources, settings and method"; ASM["A1"].font = TITLE
    lines = [("Polymarket Gamma API", f"{GAMMA}/events?tag_slug=earnings (closed and open contracts)", download_utc),
             ("Polymarket CLOB API", f"{CLOB}/prices-history, hourly YES prices", download_utc),
             ("Yahoo Finance via yfinance", "Daily adjusted close (stocks and " + MARKET + "), EPS estimates and reported EPS", download_utc)]
    head(ASM, 3, ["Source", "What", "Downloaded (UTC)"])
    for i, row in enumerate(lines):
        for j, v in enumerate(row):
            ASM.cell(row=4 + i, column=1 + j, value=v)
    rr = 8
    for lab, v in [("Reports from", START_DATE), ("Odds read at (UTC hour, report day)", CUTOFF_HOUR_UTC), ("Second reading, days earlier", LEAD_DAYS),
                   ("Minimum volume (USD)", MIN_VOLUME), ("Earlier events needed for the base rate", MIN_PRIOR_EVENTS),
                   ("Past reports needed for a company's beat rate", MIN_HISTORY), ("Market index for abnormal returns", MARKET)]:
        ASM.cell(row=rr, column=1, value=lab); c = ASM.cell(row=rr, column=2, value=v); c.font = BLUE; rr += 1
    rr += 1
    for n in ["A contract resolves YES when reported EPS exceeds the consensus estimate fixed when the contract was listed (see each contract's rules).",
              "Contracts were kept if resolved, with volume of at least the minimum, a ticker found in the question, prices on Yahoo Finance and odds before the cut-off.",
              "Yahoo Finance estimates may differ slightly from the threshold Polymarket used. They serve only for the company's past beat rate and the analyst surprise.",
              "The running base rate uses only events before each report, so it is a fair out-of-sample benchmark.",
              "Abnormal return = stock log return minus market log return over the same two-day window. No beta adjustment.",
              "Small or thin contracts can be pushed around. Results describe one sample and are associations, not causes."]:
        ASM.cell(row=rr, column=1, value="- " + n); rr += 1
    ASM.column_dimensions["A"].width = 46; ASM.column_dimensions["B"].width = 80; ASM.column_dimensions["C"].width = 18

    add_extras(wb, opts, download_utc, hist_px, edgar, short)
    add_robust(wb)
    add_cutoffs(wb)
    add_lookup(wb)
    nk = 0
    if km is not None and len(km):
        nk = build_kalshi(wb, km, kodds, prices, PE, EE, download_utc, head, ASM,
                          dict(BOLD=BOLD, NOTE=NOTE, TITLE=TITLE, SEC=SEC, BLUE=BLUE))
    if fed is not None and len(fed.get("markets", [])):
        try:
            build_fed(wb, fed, download_utc, head, dict(BOLD=BOLD, NOTE=NOTE, TITLE=TITLE, SEC=SEC, BLUE=BLUE))
        except Exception as e:
            print(f"   Fed sheets skipped: {e}", flush=True)
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = False
    wb.active = 0; wb.worksheets[0].sheet_view.tabSelected = True
    for n, col in (("Results", "1F3A5F"), ("Events", "2E7D32"), ("Live", "2E7D32"), ("Assumptions", "A6A6A6"),
                   ("Kalshi_Results", "1F3A5F"), ("Kalshi_Events", "2E7D32"), ("Fed", "1F3A5F"), ("Fed_Events", "2E7D32")):
        if n in wb.sheetnames:
            wb[n].sheet_properties.tabColor = col
    for n in wb.sheetnames:
        if n.startswith("Raw_"):
            wb[n].sheet_properties.tabColor = "C65911"
    fix_locale(wb)
    wb.save(OUTPUT_FILE)
    return len(good), nk


def add_extras(wb, opts=None, download_utc="", hist_px=None, edgar=None, short=None):
    """Extra tests on Events: run-up before and drift after the report, liquid-only check, combined forecast,
    and the option-implied move for upcoming reports. Reads everything it needs from the workbook itself."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as L
    from openpyxl.worksheet.formula import ArrayFormula
    BOLD = Font(name="Arial", bold=True); NOTE = Font(name="Arial", italic=True, color="666666")
    SEC = Font(name="Arial", bold=True, size=11, color="1F3A5F"); BLUE = Font(name="Arial", color="0000FF")
    HEAD = PatternFill("solid", fgColor="DDE4EE")

    def head(ws, r, vals, c0=1):
        for j, v in enumerate(vals):
            c = ws.cell(row=r, column=c0 + j, value=v); c.font = BOLD; c.fill = HEAD
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 45

    EVT, RES, LIV, ASM, RP, RM = (wb[n] for n in ("Events", "Results", "Live", "Assumptions", "Raw_Prices", "Raw_Markets"))
    PE = 4
    while RP.cell(row=PE + 1, column=1).value is not None:
        PE += 1
    pcol = {RP.cell(row=4, column=j).value: L(j) for j in range(2, RP.max_column + 1)}
    E_ = 4
    while EVT.cell(row=E_ + 1, column=1).value is not None:
        E_ += 1
    if E_ < 6:
        return
    # editable input for the liquid-only check
    ar = ASM.max_row + 2
    ASM.cell(row=ar, column=1, value="Liquid contract: minimum volume (USD), used on Results part F")
    c = ASM.cell(row=ar, column=2, value=LIQ_VOLUME); c.font = BLUE
    LIQ = f"Assumptions!$B${ar}"
    ASM.cell(row=ar + 1, column=1, value="Days for the run-up before and the drift after the report (trading days)")
    ASM.cell(row=ar + 1, column=2, value=DRIFT_DAYS).font = BLUE
    ASM.cell(row=ar + 2, column=1, value="Stricter liquidity tier: minimum volume (USD), Results part F")
    ASM.cell(row=ar + 2, column=2, value=LIQ_VOLUME_2).font = BLUE
    LIQ2 = f"Assumptions!$B${ar + 2}"
    ASM.cell(row=ar + 3, column=1, value=f"Change in p over {LEAD_DAYS} days counted as 'rose' or 'fell', Results part H")
    ASM.cell(row=ar + 3, column=2, value=DP_BAND).font = BLUE
    DPB = f"Assumptions!$B${ar + 3}"
    ASM.cell(row=ar + 4, column=1, value="Spread and fees per $1 prediction-market contract, Results part O")
    ASM.cell(row=ar + 4, column=2, value=TRADE_COST).font = BLUE
    COST = f"Assumptions!$B${ar + 4}"

    head(EVT, 4, [f"Run-up: abnormal return over {DRIFT_DAYS} days before the report",
                  f"Drift: abnormal return over {DRIFT_DAYS} days after the reaction window",
                  "Liquid (volume at least the Assumptions input)", "Combined forecast: average of p and company rate",
                  "Brier: combined", "Liquid and all three forecasts", f"Change in p over the last {LEAD_DAYS} days",
                  "Helper: y - p (rows with a change in p)", "Helper: p (same rows)", "EPS type (Raw_Markets)",
                  "Helper: y - p net of p", "Helper: change in p net of p",
                  f"Volatility before the report ({VOL_DAYS}-day SD of daily abnormal returns)", "Gap: p minus company record"], c0=27)
    cM = pcol.get(MARKET)
    D = DRIFT_DAYS
    for r in range(5, E_ + 1):
        raw = EVT.cell(row=r, column=1).value
        sym = yf_symbol(str(RM.cell(row=raw, column=3).value))
        cS = pcol.get(sym)
        if cS and cM:
            pre = f"MATCH(C{r}-1,Raw_Prices!$A$5:$A${PE},1)"; post = f"(MATCH(C{r},Raw_Prices!$A$5:$A${PE},1)+1)"
            S_ = f"Raw_Prices!${cS}$5:${cS}${PE}"; M_ = f"Raw_Prices!${cM}$5:${cM}${PE}"
            EVT[f"AA{r}"] = f'=IFERROR(LN(INDEX({S_},{pre})/INDEX({S_},{pre}-{D}))-LN(INDEX({M_},{pre})/INDEX({M_},{pre}-{D})),"")'
            EVT[f"AB{r}"] = f'=IFERROR(LN(INDEX({S_},{post}+{D})/INDEX({S_},{post}))-LN(INDEX({M_},{post}+{D})/INDEX({M_},{post})),"")'
        EVT[f"AC{r}"] = f"=IF(K{r}>={LIQ},1,0)"
        EVT[f"AD{r}"] = f'=IF(AND(ISNUMBER(I{r}),ISNUMBER(M{r})),(I{r}+M{r})/2,"")'
        EVT[f"AE{r}"] = f'=IF(ISNUMBER(AD{r}),(AD{r}-E{r})^2,"")'
        EVT[f"AF{r}"] = f"=R{r}*AC{r}"
        EVT[f"AG{r}"] = f'=IF(AND(ISNUMBER(I{r}),ISNUMBER(J{r})),I{r}-J{r},"")'; EVT[f"AG{r}"].number_format = "+0.000;-0.000"
        EVT[f"AJ{r}"] = f"=Raw_Markets!F{raw}"
        EVT[f"AH{r}"] = f'=IF(ISNUMBER(AG{r}),V{r},"")'; EVT[f"AI{r}"] = f'=IF(ISNUMBER(AG{r}),I{r},"")'
        for k in ("AA", "AB"):
            EVT[f"{k}{r}"].number_format = "+0.00%;-0.00%"
        EVT[f"AD{r}"].number_format = "0.000"; EVT[f"AE{r}"].number_format = "0.000"
    R = lambda c: f"Events!${c}$5:${c}${E_}"
    tt = lambda y, x: f'IFERROR(SLOPE({y},{x})/(STEYX({y},{x})/SQRT(DEVSQ({x}))),"n/a")'
    P2 = "+0.00%;-0.00%"

    # ---------- Results E: before and after ----------
    rr = RES.max_row + 3
    RES.cell(row=rr, column=1, value=f"E. Before and after the report: abnormal returns over {D} trading days").font = SEC; rr += 1
    head(RES, rr, ["Average abnormal return", "Polymarket p at least 0.8", "Polymarket p below 0.8", "Difference"]); rr += 1
    for lab, col, yv in ((f"Run-up before, company then beat", "AA", 1), ("Run-up before, company then missed", "AA", 0),
                         ("Drift after, company beat", "AB", 1), ("Drift after, company missed", "AB", 0)):
        RES.cell(row=rr, column=1, value=lab)
        RES.cell(row=rr, column=2, value=f'=IFERROR(AVERAGEIFS({R(col)},{R("E")},{yv},{R("I")},">=0.8"),"n/a")').number_format = P2
        RES.cell(row=rr, column=3, value=f'=IFERROR(AVERAGEIFS({R(col)},{R("E")},{yv},{R("I")},"<0.8"),"n/a")').number_format = P2
        RES.cell(row=rr, column=4, value=f'=IFERROR(B{rr}-C{rr},"n/a")').number_format = P2
        rr += 1
    rr += 1
    head(RES, rr, ["Regression", "Slope", "t", "R squared", "Events"]); rr += 1
    for lab, y, x in (("Run-up before on Polymarket p (was a likely beat already in the price?)", "AA", "I"),
                      ("Drift after on Polymarket surprise y - p", "AB", "V"),
                      ("Drift after on analyst surprise (%)", "AB", "N"),
                      ("Drift after on the reaction itself (continuation if positive)", "AB", "U")):
        RES.cell(row=rr, column=1, value=lab)
        RES.cell(row=rr, column=2, value=f'=IFERROR(SLOPE({R(y)},{R(x)}),"n/a")').number_format = "0.0000"
        RES.cell(row=rr, column=3, value="=" + tt(R(y), R(x))).number_format = "0.0"
        RES.cell(row=rr, column=4, value=f'=IFERROR(RSQ({R(y)},{R(x)}),"n/a")').number_format = "0.000"
        RES.cell(row=rr, column=5, value=f'=SUMPRODUCT(--ISNUMBER({R(y)}),--ISNUMBER({R(x)}))').number_format = "0"
        rr += 1

    # ---------- Results F: liquidity tiers ----------
    def group_table(rr, title, groups):
        """Key results for several sub-samples. groups = (label, SUMPRODUCT mask or None, AVERAGEIFS extra criteria)."""
        RES.cell(row=rr, column=1, value=title).font = SEC; rr += 1
        head(RES, rr, [""] + [g[0] for g in groups]); rr += 1
        t0 = rr
        labels = ["Events", "Share of all events", "Median contract volume (USD)", "Events with all three forecasts",
                  "Brier, Polymarket", "Brier, running base rate", "Brier, company past beat rate", "Skill vs base rate",
                  "Skill vs company history", "Correct side", "Average abnormal return: beat with p at least 0.8",
                  "Average abnormal return: miss with p at least 0.8", "Misses with p at least 0.8 (count)"]
        for k, lab in enumerate(labels):
            RES.cell(row=t0 + k, column=1, value=lab)
        for j, (lab, m, crit) in enumerate(groups):
            c = L(2 + j); mm = f",{m}" if m else ""
            ev_ = f"=SUMPRODUCT(--ISNUMBER({R('E')}){mm})"
            fm = [ev_, f"={c}{t0}/COUNT({R('E')})",
                  (f"=MEDIAN({R('K')})" if not m else ArrayFormula(f"{c}{t0 + 2}", f"=IFERROR(MEDIAN(IF({m[3:-1]},{R('K')})),\"n/a\")")),
                  f"=SUMPRODUCT({R('R')}{mm})"]
            for col in ("O", "P", "Q"):
                fm.append(f'=IFERROR(SUMPRODUCT({R("R")},{R(col)}{mm})/{c}{t0 + 3},"n/a")')
            fm += [f'=IFERROR(1-{c}{t0 + 4}/{c}{t0 + 5},"n/a")', f'=IFERROR(1-{c}{t0 + 4}/{c}{t0 + 6},"n/a")',
                   f'=IFERROR(SUMPRODUCT({R("Y")}{mm})/SUMPRODUCT(--ISNUMBER({R("Y")}){mm}),"n/a")',
                   f'=IFERROR(AVERAGEIFS({R("U")},{R("E")},1,{R("I")},">=0.8"{crit}),"n/a")',
                   f'=IFERROR(AVERAGEIFS({R("U")},{R("E")},0,{R("I")},">=0.8"{crit}),"n/a")',
                   f'=COUNTIFS({R("E")},0,{R("I")},">=0.8"{crit})']
            nfs = ["0", "0.0%", "#,##0", "0", "0.000", "0.000", "0.000", "0.0%", "0.0%", "0.0%", P2, P2, "0"]
            for k, (f_, nf) in enumerate(zip(fm, nfs)):
                cell = RES.cell(row=t0 + k, column=2 + j, value=f_); cell.number_format = nf
        return t0 + len(labels)

    rr += 1
    rr = group_table(rr, "F. Robustness: liquidity tiers (thresholds on Assumptions)",
                     [("All events", None, ""),
                      ("Volume at least tier 1", f"--({R('K')}>={LIQ})", f',{R("K")},">="&{LIQ}'),
                      ("Volume at least tier 2", f"--({R('K')}>={LIQ2})", f',{R("K")},">="&{LIQ2}')])

    # ---------- Results G: combined forecast ----------
    rr += 1
    RES.cell(row=rr, column=1, value="G. Does combining sources help? Simple average of Polymarket p and the company's beat rate").font = SEC; rr += 1
    head(RES, rr, ["Same events (all three forecasts)", "Brier", "Skill of the combination vs this source"]); rr += 1
    g0 = rr
    for lab, col in (("Polymarket alone", "O"), ("Company history alone", "Q"), ("Combined (average)", "AE")):
        RES.cell(row=rr, column=1, value=lab)
        RES.cell(row=rr, column=2, value=f"=SUMPRODUCT({R('R')},{R(col)})/SUM({R('R')})").number_format = "0.000"
        if col != "AE":
            RES.cell(row=rr, column=3, value=f'=IFERROR(1-B{g0 + 2}/B{rr},"n/a")').number_format = "0.0%"
        rr += 1
    RES.cell(row=rr, column=1, value="A positive skill means the average beats that source alone. Equal weights, no fitting, so no look-ahead.").font = NOTE

    # ---------- Results H: change in the odds over the last week ----------
    rr += 2
    RES.cell(row=rr, column=1, value=f"H. Did the odds move in the last {LEAD_DAYS} days? (change in p, band on Assumptions)").font = SEC; rr += 1
    head(RES, rr, ["", "Events", "Average change in p", "Average p at cut-off", "Share that beat", "Beat share minus p",
                   "Average run-up before", "Average reaction"]); rr += 1
    G_ = R("AG")
    for lab, crit in ((f"p fell", f'"<="&-{DPB}'), ("p roughly flat", None), ("p rose", f'">="&{DPB}')):
        RES.cell(row=rr, column=1, value=lab)
        cr = f"{G_},{crit}" if crit else f'{G_},">"&-{DPB},{G_},"<"&{DPB}'
        RES.cell(row=rr, column=2, value=f"=COUNTIFS({cr})")
        for j, col in enumerate(("AG", "I", "E", None, "AA", "U")):
            if col is None:
                fm = f'=IFERROR(E{rr}-D{rr},"n/a")'
            else:
                fm = f'=IFERROR(AVERAGEIFS({R(col)},{cr}),"n/a")'
            nf = {"AG": "+0.000;-0.000", "I": "0.0%", "E": "0.0%", None: "+0.0%;-0.0%", "AA": P2, "U": P2}[col]
            RES.cell(row=rr, column=3 + j, value=fm).number_format = nf
        rr += 1
    head(RES, rr, ["Regression", "Slope", "t", "R squared", "Events"]); rr += 1
    for lab, y, x in (("Outcome minus p (y - p) on the change in p: does a rising p under-react?", "V", "AG"),
                      ("Run-up before on the change in p", "AA", "AG"),
                      ("Reaction on the change in p", "U", "AG")):
        RES.cell(row=rr, column=1, value=lab)
        RES.cell(row=rr, column=2, value=f'=IFERROR(SLOPE({R(y)},{R(x)}),"n/a")').number_format = "0.0000"
        RES.cell(row=rr, column=3, value="=" + tt(R(y), R(x))).number_format = "0.0"
        RES.cell(row=rr, column=4, value=f'=IFERROR(RSQ({R(y)},{R(x)}),"n/a")').number_format = "0.000"
        RES.cell(row=rr, column=5, value=f'=SUMPRODUCT(--ISNUMBER({R(y)}),--ISNUMBER({R(x)}))').number_format = "0"
        rr += 1
    h0 = rr
    hl = [("Outcome minus p on the change in p, holding the level of p fixed: slope", f'=IFERROR(SLOPE({R("AK")},{R("AL")}),"n/a")', "0.0000"),
          ("t statistic of that slope", f'=IFERROR(B{h0}/SQRT((SUMSQ({R("AK")})-B{h0}^2*DEVSQ({R("AL")}))/(COUNT({R("AK")})-3)/DEVSQ({R("AL")})),"n/a")', "0.0"),
          ("Helper: intercept, y - p on p", f'=IFERROR(INTERCEPT({R("AH")},{R("AI")}),0)', "0.0000"),
          ("Helper: slope, y - p on p", f'=IFERROR(SLOPE({R("AH")},{R("AI")}),0)', "0.0000"),
          ("Helper: intercept, change in p on p", f'=IFERROR(INTERCEPT({R("AG")},{R("AI")}),0)', "0.0000"),
          ("Helper: slope, change in p on p", f'=IFERROR(SLOPE({R("AG")},{R("AI")}),0)', "0.0000")]
    for k, (lab, fm, nf) in enumerate(hl):
        RES.cell(row=h0 + k, column=1, value=lab); RES.cell(row=h0 + k, column=2, value=fm).number_format = nf
    for r in range(5, E_ + 1):
        EVT[f"AK{r}"] = f'=IF(ISNUMBER(AG{r}),AH{r}-(Results!$B${h0 + 2}+Results!$B${h0 + 3}*AI{r}),"")'
        EVT[f"AL{r}"] = f'=IF(ISNUMBER(AG{r}),AG{r}-(Results!$B${h0 + 4}+Results!$B${h0 + 5}*AI{r}),"")'
    rr = h0 + len(hl)
    RES.cell(row=rr, column=1, value=("Reading: if 'beat share minus p' is positive when p fell, the crowd became too pessimistic "
                                      "(odds over-reacted). Near zero means the final p already uses the information. "
                                      "The controlled slope removes any effect of the level of p (rows above use the two-step method)."))
    RES.cell(row=rr, column=1).font = NOTE

    # ---------- Results J: GAAP vs non-GAAP ----------
    rr += 2
    rr = group_table(rr, "J. GAAP and non-GAAP contracts separately",
                     [("GAAP", f'--({R("AJ")}="GAAP")', f',{R("AJ")},"GAAP"'),
                      ("non-GAAP", f'--({R("AJ")}="non-GAAP")', f',{R("AJ")},"non-GAAP"')])

    # ---------- Events: volatility before the report and the crowd-vs-record gap ----------
    V_ = VOL_DAYS
    for r in range(5, E_ + 1):
        raw = EVT.cell(row=r, column=1).value
        cS = pcol.get(yf_symbol(str(RM.cell(row=raw, column=3).value)))
        if cS and cM:
            pre = f"MATCH(C{r}-1,Raw_Prices!$A$5:$A${PE},1)"
            S_ = f"Raw_Prices!${cS}$5:${cS}${PE}"; M_ = f"Raw_Prices!${cM}$5:${cM}${PE}"
            rng = lambda X, a, b: f"INDEX({X},{pre}-{a}):INDEX({X},{pre}-{b})"
            EVT[f"AM{r}"] = ArrayFormula(f"AM{r}", (f'=IFERROR(STDEV(LN({rng(S_, V_ - 1, 0)}/{rng(S_, V_, 1)})'
                                                    f'-LN({rng(M_, V_ - 1, 0)}/{rng(M_, V_, 1)})),"")'))
            EVT[f"AM{r}"].number_format = "0.00%"
        EVT[f"AN{r}"] = f'=IF(AND(ISNUMBER(I{r}),ISNUMBER(M{r})),I{r}-M{r},"")'; EVT[f"AN{r}"].number_format = "+0.00;-0.00"

    # ---------- Results K: volatility and the size of the move ----------
    rr += 2
    RES.cell(row=rr, column=1, value=f"K. Is the size of the move predictable? |Abnormal return| on volatility before the report").font = SEC; rr += 1
    head(RES, rr, ["", "Intercept", "Slope", "t", "R squared", "Events"]); rr += 1
    RES.cell(row=rr, column=1, value=f"|Abnormal return| = a + b x {VOL_DAYS}-day volatility")
    RES.cell(row=rr, column=2, value=f'=IFERROR(INTERCEPT({R("X")},{R("AM")}),"n/a")').number_format = "0.0000"
    RES.cell(row=rr, column=3, value=f'=IFERROR(SLOPE({R("X")},{R("AM")}),"n/a")').number_format = "0.000"
    RES.cell(row=rr, column=4, value="=" + tt(R("X"), R("AM"))).number_format = "0.0"
    RES.cell(row=rr, column=5, value=f'=IFERROR(RSQ({R("X")},{R("AM")}),"n/a")').number_format = "0.000"
    RES.cell(row=rr, column=6, value=f'=SUMPRODUCT(--ISNUMBER({R("X")}),--ISNUMBER({R("AM")}))')
    VOLROW = rr; rr += 2
    head(RES, rr, ["Volatility before the report", "From", "To", "Events", "Average |abnormal return|", "Fell 10% or more", "Rose 10% or more"]); rr += 1
    for k, (lab, lo, hi) in enumerate((("Lowest third", 0, 1 / 3), ("Middle third", 1 / 3, 2 / 3), ("Highest third", 2 / 3, 1))):
        RES.cell(row=rr, column=1, value=lab)
        RES.cell(row=rr, column=2, value=f"=PERCENTILE({R('AM')},{lo})").number_format = "0.00%"
        RES.cell(row=rr, column=3, value=f"=PERCENTILE({R('AM')},{hi})+{'0.000001' if hi == 1 else '0'}").number_format = "0.00%"
        cr = f'{R("AM")},">="&B{rr},{R("AM")},"<"&C{rr}'
        RES.cell(row=rr, column=4, value=f"=COUNTIFS({cr})")
        RES.cell(row=rr, column=5, value=f'=IFERROR(AVERAGEIFS({R("X")},{cr}),"n/a")').number_format = "0.0%"
        RES.cell(row=rr, column=6, value=f'=IFERROR(COUNTIFS({cr},{R("U")},"<=-0.1")/COUNTIFS({cr},{R("U")},">=-10"),"n/a")').number_format = "0.0%"
        RES.cell(row=rr, column=7, value=f'=IFERROR(COUNTIFS({cr},{R("U")},">=0.1")/COUNTIFS({cr},{R("U")},">=-10"),"n/a")').number_format = "0.0%"
        rr += 1
    RES.cell(row=rr, column=1, value=("Reading: volatility before the report goes with the size of the move in both directions, "
                                      "not with its direction. The Live sheet uses this line for the expected move.")).font = NOTE

    # ---------- Results L: crowd vs the company's record ----------
    rr += 2
    RES.cell(row=rr, column=1, value=f"L. When the crowd and the company's record disagree, who is right? (gap of {GAP} or more)").font = SEC; rr += 1
    head(RES, rr, ["", "Events", "Average crowd p", "Average company record", "Share that beat", "Average abnormal return"]); rr += 1
    DIS = {}
    for key, lab, cr in (("low", "Crowd much less confident than the record", f'{R("AN")},"<="&-{GAP}'),
                         ("mid", "Roughly agree", f'{R("AN")},">"&-{GAP},{R("AN")},"<"&{GAP}'),
                         ("high", "Crowd much more confident than the record", f'{R("AN")},">="&{GAP}')):
        RES.cell(row=rr, column=1, value=lab)
        RES.cell(row=rr, column=2, value=f"=COUNTIFS({cr})")
        for j, col in enumerate(("I", "M", "E", "U")):
            RES.cell(row=rr, column=3 + j, value=f'=IFERROR(AVERAGEIFS({R(col)},{cr}),"n/a")').number_format = P2 if col == "U" else "0.0%"
        DIS[key] = rr; rr += 1
    RES.cell(row=rr, column=1, value=("Reading: compare the share that beat with each forecast. The closer one was right more often. "
                                      "The company record is the analysts' track record: how often the company beat their consensus.")).font = NOTE

    # ---------- Live: option-implied move ----------
    if opts is not None:
        RO = wb.create_sheet("Raw_Options")
        RO["A1"] = (f"At-the-money straddle on the first expiry after the report date, Yahoo Finance via yfinance option_chain. "
                    f"Snapshot at download {download_utc} UTC. Mid price, last trade if no quote."); RO["A1"].font = BOLD
        head(RO, 4, ["Ticker", "Expiry", "Spot", "Strike", "Call (mid)", "Put (mid)"])
        for i, o in enumerate(opts):
            if o:
                vals = [o["ticker"], o["expiry"].to_pydatetime(), o["spot"], o["strike"], o["call"], o["put"]]
                for j, v in enumerate(vals):
                    RO.cell(row=5 + i, column=1 + j, value=v)
                RO.cell(row=5 + i, column=2).number_format = "yyyy-mm-dd"
        RO.sheet_properties.tabColor = "C65911"
        # past earnings moves: long price history for these companies, reactions on every Yahoo earnings date
        RE_ = wb["Raw_Estimates"]
        EE2 = 4
        while RE_.cell(row=EE2 + 1, column=1).value is not None:
            EE2 += 1
        src = "Events"
        if hist_px is not None and len(hist_px) and EE2 >= 5:
            RH = wb.create_sheet("Raw_Hist_Prices")
            RH["A1"] = (f"Daily adjusted close from {HIST_START}, upcoming-report companies and {MARKET} only, Yahoo Finance via yfinance. "
                        f"Downloaded {download_utc} UTC. Used for each company's past earnings moves (Raw_Estimates column G)."); RH["A1"].font = BOLD
            head(RH, 4, ["Date"] + list(hist_px.columns))
            for i, (d, row) in enumerate(hist_px.iterrows()):
                RH.cell(row=5 + i, column=1, value=d.to_pydatetime()).number_format = "yyyy-mm-dd"
                for j, c in enumerate(hist_px.columns):
                    if pd.notna(row[c]):
                        RH.cell(row=5 + i, column=2 + j, value=float(row[c]))
            RH.freeze_panes = "A5"; RH.sheet_properties.tabColor = "C65911"
            HE = 4 + len(hist_px); hcol = {c: L(2 + j) for j, c in enumerate(hist_px.columns)}
            head(RE_, 4, ["|Abnormal move| on this report (live companies only, Raw_Hist_Prices)"], c0=7)
            hm = hcol.get(MARKET)
            for rr_ in range(5, EE2 + 1):
                sc = hcol.get(yf_symbol(str(RE_.cell(row=rr_, column=1).value)))
                if sc and hm:
                    pre = f"MATCH(B{rr_}-1,Raw_Hist_Prices!$A$5:$A${HE},1)"; post = f"(MATCH(B{rr_},Raw_Hist_Prices!$A$5:$A${HE},1)+1)"
                    S_ = f"Raw_Hist_Prices!${sc}$5:${sc}${HE}"; M_ = f"Raw_Hist_Prices!${hm}$5:${hm}${HE}"
                    RE_[f"G{rr_}"] = (f'=IFERROR(ABS(LN(INDEX({S_},{post})/INDEX({S_},{pre}))-LN(INDEX({M_},{post})/INDEX({M_},{pre}))),"")')
                    RE_[f"G{rr_}"].number_format = "0.00%"
            src = "Raw_Estimates"
        head(LIV, 4, ["Option-implied move (straddle / spot)", "Option expiry", "Company's average past |abnormal move|",
                      "Implied / past", "Options reading", "Past reports used"], c0=11)
        n_live = 0
        while LIV.cell(row=5 + n_live, column=1).value is not None:
            n_live += 1
        for i in range(n_live):
            r = 5 + i
            LIV[f"K{r}"] = f'=IFERROR((Raw_Options!E{r}+Raw_Options!F{r})/Raw_Options!C{r},"n/a")'; LIV[f"K{r}"].number_format = "0.0%"
            LIV[f"L{r}"] = f'=IF(ISNUMBER(Raw_Options!B{r}),Raw_Options!B{r},"")'; LIV[f"L{r}"].number_format = "yyyy-mm-dd"
            if src == "Raw_Estimates":
                LIV[f"M{r}"] = f'=IFERROR(AVERAGEIFS(Raw_Estimates!$G$5:$G${EE2},Raw_Estimates!$A$5:$A${EE2},A{r}),"n/a")'
                LIV[f"P{r}"] = f'=COUNTIFS(Raw_Estimates!$A$5:$A${EE2},A{r},Raw_Estimates!$G$5:$G${EE2},">=0")'
            else:
                LIV[f"M{r}"] = f'=IFERROR(AVERAGEIF({R("B")},A{r},{R("X")}),"n/a")'
                LIV[f"P{r}"] = f'=COUNTIFS({R("B")},A{r},{R("X")},">=0")'
            LIV[f"M{r}"].number_format = "0.0%"
            LIV[f"N{r}"] = f'=IF(AND(ISNUMBER(K{r}),ISNUMBER(M{r})),K{r}/M{r},"n/a")'; LIV[f"N{r}"].number_format = "0.00"
            LIV[f"O{r}"] = (f'=IF(NOT(ISNUMBER(N{r})),"no comparison",IF(P{r}<{MIN_PAST_MOVES},"too few past reports",'
                            f'IF(N{r}>=1.5,"options price a bigger move than usual",'
                            f'IF(N{r}<=0.67,"options price a smaller move than usual","in line with past moves"))))')
        for c, w in zip("KLMNOP", (13, 12, 14, 10, 34, 10)):
            LIV.column_dimensions[c].width = w
        LIV["A3"] = ("Options columns: the straddle runs to the expiry, so it includes a few non-event days, and option prices usually "
                     "exceed realised moves (volatility risk premium). Past moves: the company's reactions on its past earnings dates since "
                     f"{HIST_START}. Readings need at least {MIN_PAST_MOVES} past reports.")
        LIV["A3"].font = NOTE

    # ---------- Raw_Edgar, Raw_Short and their Events columns ----------
    head(EVT, 4, ["EDGAR row (set by the script)", "Release time (UTC, EDGAR acceptance)", "Release session",
                  "Odds read before the release?", "Odds clearly before release (1/0)", "Short interest row (set by the script)",
                  "Days to cover (latest FINRA before the report)", "Helper: (y - p) squared"], c0=41)
    ev_dates = {}
    for r in range(5, E_ + 1):
        raw = EVT.cell(row=r, column=1).value
        ev_dates[r] = (str(RM.cell(row=raw, column=3).value).upper(), pd.Timestamp(RM.cell(row=raw, column=4).value))
        EVT[f"AV{r}"] = f'=IF(ISNUMBER(V{r}),V{r}^2,"")'
    if edgar is not None and len(edgar):
        RE2 = wb.create_sheet("Raw_Edgar")
        RE2["A1"] = ("Earnings press-release filings (form 8-K, item 2.02) as downloaded: SEC EDGAR data.sec.gov/submissions. "
                     f"Downloaded {download_utc} UTC. EDGAR marks times with Z, but some are New York wall time; column G says how each was read.")
        RE2["A1"].font = BOLD
        head(RE2, 4, ["Ticker", "Filing date", "Acceptance time as published", "Release time (UTC)", "Release time (New York)",
                      "Session", "Time read as"])
        idx = {}
        for i, x in enumerate(edgar.itertuples()):
            rr_ = 5 + i
            for j, v in enumerate([x.ticker, pd.Timestamp(x.filed).to_pydatetime(), str(x.acceptance), x.release_utc.to_pydatetime(),
                                   x.release_et.to_pydatetime(), x.session, x.read_as]):
                RE2.cell(row=rr_, column=1 + j, value=v)
            RE2.cell(row=rr_, column=2).number_format = "yyyy-mm-dd"
            RE2.cell(row=rr_, column=4).number_format = "yyyy-mm-dd hh:mm"; RE2.cell(row=rr_, column=5).number_format = "yyyy-mm-dd hh:mm"
            idx.setdefault(str(x.ticker).upper(), []).append((pd.Timestamp(x.release_et).normalize(), rr_))
        RE2.freeze_panes = "A5"; RE2.sheet_properties.tabColor = "C65911"
        for r, (t, d) in ev_dates.items():
            c = [(abs((dd - d).days), rr_) for dd, rr_ in idx.get(t, []) if abs((dd - d).days) <= 1]
            if c:
                EVT[f"AO{r}"] = min(c)[1]
                EVT[f"AP{r}"] = f"=INDEX(Raw_Edgar!$D:$D,AO{r})"; EVT[f"AP{r}"].number_format = "yyyy-mm-dd hh:mm"
                EVT[f"AQ{r}"] = f"=INDEX(Raw_Edgar!$F:$F,AO{r})"
            EVT[f"AR{r}"] = (f'=IF(ISNUMBER(AP{r}),IF((AP{r}-DATE(1970,1,1))*86400>H{r}+3600,"yes","possibly after the release"),'
                             f'"no EDGAR match")')
            EVT[f"AS{r}"] = f'=IF(AR{r}="yes",1,0)'
    sidx = {}
    if short is not None and len(short):
        RS = wb.create_sheet("Raw_Short")
        RS["A1"] = ("Short interest as downloaded: FINRA Query API, otcMarket/consolidatedShortInterest (public). "
                    f"Downloaded {download_utc} UTC. Published about 8 business days after the settlement date.")
        RS["A1"].font = BOLD
        head(RS, 4, ["Symbol", "Settlement date", "Short position (shares)", "Average daily volume", "Days to cover"])
        for i, x in enumerate(short.itertuples()):
            rr_ = 5 + i
            for j, v in enumerate([x.symbol, x.settle.to_pydatetime(), x.short, x.adv, x.dtc]):
                RS.cell(row=rr_, column=1 + j, value=None if (isinstance(v, float) and np.isnan(v)) else v)
            RS.cell(row=rr_, column=2).number_format = "yyyy-mm-dd"
            sidx.setdefault(str(x.symbol).upper(), []).append((x.settle, rr_))
        RS.freeze_panes = "A5"; RS.sheet_properties.tabColor = "C65911"
        for r, (t, d) in ev_dates.items():
            c = [rr_ for dd, rr_ in sidx.get(t, []) if dd <= d - pd.Timedelta(days=SI_LAG_DAYS)]
            if c:
                EVT[f"AT{r}"] = c[-1]
                EVT[f"AU{r}"] = f'=IF(ISNUMBER(INDEX(Raw_Short!$E:$E,AT{r})),INDEX(Raw_Short!$E:$E,AT{r}),"")'
                EVT[f"AU{r}"].number_format = "0.0"

    # ---------- Results M: was the crowd's probability read before the news? ----------
    if edgar is not None and len(edgar):
        rr += 3
        RES.cell(row=rr, column=1, value=("Check on timing: share of events matched to an EDGAR earnings filing")).font = NOTE
        RES.cell(row=rr, column=2, value=f"=COUNT({R('AP')})/COUNT({R('E')})").number_format = "0.0%"
        RES.cell(row=rr, column=3, value="odds clearly read before the release:")
        RES.cell(row=rr, column=4, value=f"=SUM({R('AS')})/COUNT({R('E')})").number_format = "0.0%"
        rr += 1
        rr = group_table(rr, "M. Timing check with SEC EDGAR release times",
                         [("All events", None, ""),
                          ("Odds clearly read before the release", f"--({R('AS')}=1)", f',{R("AS")},1'),
                          ("Released before the open", f'--({R("AQ")}="before open")', f',{R("AQ")},"before open"'),
                          ("Released after the close", f'--({R("AQ")}="after close")', f',{R("AQ")},"after close"')])
        RES.cell(row=rr, column=1, value=("Reading: if the results hold for events where the odds were clearly read before the filing, "
                                          "the crowd's edge does not come from trading after the news.")).font = NOTE

    # ---------- Results N: short interest ----------
    if short is not None and len(short):
        rr += 3
        RES.cell(row=rr, column=1, value="N. Do heavily shorted stocks react differently? (days to cover, latest FINRA figure before the report)").font = SEC; rr += 1
        RES.cell(row=rr, column=1, value="Top third starts at (days to cover)")
        RES.cell(row=rr, column=2, value=f"=PERCENTILE({R('AU')},2/3)").number_format = "0.0"
        THR = f"Results!$B${rr}"; thr_row = rr; rr += 1
        head(RES, rr, ["Average abnormal return", "Most shorted third", "Other stocks", "Difference", "Events (most shorted)", "Events (other)"]); rr += 1
        for lab, yv, pc in (("Beat, p at least 0.8", 1, '">=0.8"'), ("Miss, p at least 0.8", 0, '">=0.8"'),
                            ("Beat, p below 0.8", 1, '"<0.8"'), ("Miss, p below 0.8", 0, '"<0.8"')):
            base_cr = f'{R("E")},{yv},{R("I")},{pc}'
            RES.cell(row=rr, column=1, value=lab)
            RES.cell(row=rr, column=2, value=f'=IFERROR(AVERAGEIFS({R("U")},{base_cr},{R("AU")},">="&$B${thr_row}),"n/a")').number_format = P2
            RES.cell(row=rr, column=3, value=f'=IFERROR(AVERAGEIFS({R("U")},{base_cr},{R("AU")},"<"&$B${thr_row}),"n/a")').number_format = P2
            RES.cell(row=rr, column=4, value=f'=IFERROR(B{rr}-C{rr},"n/a")').number_format = P2
            RES.cell(row=rr, column=5, value=f'=COUNTIFS({base_cr},{R("AU")},">="&$B${thr_row},{R("U")},">=-10")')
            RES.cell(row=rr, column=6, value=f'=COUNTIFS({base_cr},{R("AU")},"<"&$B${thr_row},{R("U")},">=-10")')
            rr += 1
        RES.cell(row=rr, column=1, value="Fell 10% or more on the report")
        RES.cell(row=rr, column=2, value=f'=IFERROR(COUNTIFS({R("AU")},">="&$B${thr_row},{R("U")},"<=-0.1")/COUNTIFS({R("AU")},">="&$B${thr_row},{R("U")},">=-10"),"n/a")').number_format = "0.0%"
        RES.cell(row=rr, column=3, value=f'=IFERROR(COUNTIFS({R("AU")},"<"&$B${thr_row},{R("U")},"<=-0.1")/COUNTIFS({R("AU")},"<"&$B${thr_row},{R("U")},">=-10"),"n/a")').number_format = "0.0%"
        rr += 1
        RES.cell(row=rr, column=1, value=("Reading: short sellers are positioned for bad news. A larger fall on a priced-in miss among the most "
                                          "shorted stocks would point to crowded positioning, a risk hedge funds watch.")).font = NOTE

    # ---------- Results O: simple trading rules on the prediction market ----------
    rr += 3
    RES.cell(row=rr, column=1, value="O. Simple trading rules on the prediction market itself ($1 contracts, held to resolution)").font = SEC; rr += 1
    RES.cell(row=rr, column=1, value="Split date for the two halves (median report date)")
    RES.cell(row=rr, column=2, value=f"=MEDIAN({R('C')})").number_format = "yyyy-mm-dd"; MID = f"$B${rr}"; rr += 1
    head(RES, rr, ["Rule", "Bets", "Average price paid", "Win rate", "Average gain per contract", "Return per $ before costs",
                   "t statistic", "Return per $ after costs", "Return, first half", "Return, second half"]); rr += 1
    rules = (("Buy YES when p is below 0.6", f'{R("I")},"<0.6"', 1),
             ("Buy YES when p fell over the last week (band on Assumptions)", f'{R("AG")},"<="&-{DPB}', 1),
             ("Buy YES when the crowd doubts the company's record", f'{R("AN")},"<="&-{GAP}', 1),
             ("Buy NO when p is at least 0.8", f'{R("I")},">=0.8"', -1))
    for lab, cr, sg in rules:
        RES.cell(row=rr, column=1, value=lab)
        n_ = f'COUNTIFS({cr},{R("V")},">=-2")'
        mean_ = f'AVERAGEIFS({R("V")},{cr})'
        price_ = f'AVERAGEIFS({R("I")},{cr})' if sg == 1 else f'(1-AVERAGEIFS({R("I")},{cr}))'
        RES.cell(row=rr, column=2, value=f"={n_}")
        RES.cell(row=rr, column=3, value=f'=IFERROR({price_},"n/a")').number_format = "0.00"
        RES.cell(row=rr, column=4, value=f'=IFERROR({"" if sg == 1 else "1-"}AVERAGEIFS({R("E")},{cr}),"n/a")').number_format = "0.0%"
        RES.cell(row=rr, column=5, value=f'=IFERROR({sg}*{mean_},"n/a")').number_format = "+0.000;-0.000"
        RES.cell(row=rr, column=6, value=f'=IFERROR(E{rr}/C{rr},"n/a")').number_format = "+0.0%;-0.0%"
        RES.cell(row=rr, column=7, value=(f'=IFERROR(E{rr}/(SQRT((SUMIFS({R("AV")},{cr})-B{rr}*{mean_}^2)/(B{rr}-1))/SQRT(B{rr})),"n/a")')).number_format = "0.0"
        RES.cell(row=rr, column=8, value=f'=IFERROR((E{rr}-{COST})/C{rr},"n/a")').number_format = "+0.0%;-0.0%"
        for j, op in enumerate(('"<"&' + MID, '">="&' + MID)):
            pr_h = (f'AVERAGEIFS({R("I")},{cr},{R("C")},{op})' if sg == 1 else f'(1-AVERAGEIFS({R("I")},{cr},{R("C")},{op}))')
            RES.cell(row=rr, column=9 + j, value=f'=IFERROR({sg}*AVERAGEIFS({R("V")},{cr},{R("C")},{op})/{pr_h},"n/a")').number_format = "+0.0%;-0.0%"
        rr += 1
    RES.cell(row=rr, column=1, value=("Reading: a rule is only interesting if it earns more than its costs in both halves. Several rules are tested, "
                                      "so one good result can be luck. Treat these as ideas to test on new reports, not as a strategy.")).font = NOTE

    # ---------- Live: scenario card ----------
    n_live = 0
    while LIV.cell(row=5 + n_live, column=1).value is not None:
        n_live += 1
    if n_live:
        RL_ = wb["Raw_Live"]
        if "Raw_Hist_Prices" in wb.sheetnames:
            SRC = wb["Raw_Hist_Prices"]; sname = "Raw_Hist_Prices"
        else:
            SRC = RP; sname = "Raw_Prices"
        SL = 4
        while SRC.cell(row=SL + 1, column=1).value is not None:
            SL += 1
        scol = {SRC.cell(row=4, column=j).value: L(j) for j in range(2, SRC.max_column + 1)}
        sm = scol.get(MARKET)
        head(LIV, 4, [f"Volatility, last {VOL_DAYS} days", "Expected size of the move", "Past move if it beats (same confidence)",
                      "Past move if it misses (same confidence)", "Chance of a miss (1 - p)", "Expected cost of a surprise miss",
                      "Market liquidity", "Past beat rate in this crowd-vs-record group", "Scenario reading"], c0=17)
        hiP = HIGH_P
        for i in range(n_live):
            r = 5 + i
            sc = scol.get(yf_symbol(str(RL_.cell(row=5 + i, column=3).value)))
            if sc and sm and SL - V_ > 5:
                a, b = SL - V_ + 1, SL
                LIV[f"Q{r}"] = ArrayFormula(f"Q{r}", (f'=IFERROR(STDEV(LN({sname}!{sc}{a}:{sc}{b}/{sname}!{sc}{a - 1}:{sc}{b - 1})'
                                                      f'-LN({sname}!{sm}{a}:{sm}{b}/{sname}!{sm}{a - 1}:{sm}{b - 1})),"n/a")'))
            else:
                LIV[f"Q{r}"] = "n/a"
            LIV[f"Q{r}"].number_format = "0.00%"
            LIV[f"R{r}"] = f'=IFERROR(Results!$B${VOLROW}+Results!$C${VOLROW}*Q{r},"n/a")'; LIV[f"R{r}"].number_format = "0.0%"
            for col, yv in (("S", 1), ("T", 0)):
                LIV[f"{col}{r}"] = (f'=IFERROR(IF(F{r}>={hiP},AVERAGEIFS({R("U")},{R("E")},{yv},{R("I")},">={hiP}"),'
                                    f'AVERAGEIFS({R("U")},{R("E")},{yv},{R("I")},"<{hiP}")),"n/a")')
                LIV[f"{col}{r}"].number_format = "+0.0%;-0.0%"
            LIV[f"U{r}"] = f"=1-F{r}"; LIV[f"U{r}"].number_format = "0%"
            LIV[f"V{r}"] = f'=IF(ISNUMBER(T{r}),U{r}*T{r},"n/a")'; LIV[f"V{r}"].number_format = "+0.0%;-0.0%"
            LIV[f"W{r}"] = f'=IF(E{r}>={LIQ},"ok","volume still low, read with care")'
            LIV[f"X{r}"] = (f'=IF(NOT(ISNUMBER(H{r})),"n/a",IF(H{r}<=-{GAP},Results!$E${DIS["low"]},'
                            f'IF(H{r}>={GAP},Results!$E${DIS["high"]},Results!$E${DIS["mid"]})))')
            LIV[f"X{r}"].number_format = "0%"
            LIV[f"Y{r}"] = (f'=IF(AND(ISNUMBER(H{r}),H{r}<=-{GAP}),"crowd doubts a company that usually beats: watch for a miss",'
                            f'IF(AND(F{r}>={hiP},ISNUMBER(T{r}),T{r}<=-0.05),"confident crowd: little upside on a beat, large downside on a miss",'
                            f'IF(AND(ISNUMBER(Q{r}),Q{r}>=Results!$B${VOLROW + 5}),"volatile stock: expect a large move either way",'
                            f'"no strong signal")))')
        for c, w in zip("QRSTUVWXY", (11, 11, 12, 12, 10, 12, 16, 13, 52)):
            LIV.column_dimensions[c].width = w
        if sidx:
            head(LIV, 4, ["Days to cover (latest FINRA)", "Heavily shorted?"], c0=26)
            for i in range(n_live):
                r = 5 + i
                t = str(RL_.cell(row=5 + i, column=3).value).upper()
                c = sidx.get(t, [])
                if c:
                    LIV[f"Z{r}"] = f'=IF(ISNUMBER(INDEX(Raw_Short!$E:$E,{c[-1][1]})),INDEX(Raw_Short!$E:$E,{c[-1][1]}),"n/a")'
                    LIV[f"Z{r}"].number_format = "0.0"
                    LIV[f"AA{r}"] = (f'=IF(AND(ISNUMBER(Z{r}),ISNUMBER({THR})),IF(Z{r}>={THR},"yes","no"),"n/a")'
                                     if short is not None and len(short) else '"n/a"')
                else:
                    LIV[f"Z{r}"] = "n/a"
            LIV.column_dimensions["Z"].width = 11; LIV.column_dimensions["AA"].width = 24


def add_robust(wb, beta_days=100, gap_days=10, min_train=250):
    """Results part Q: robustness of the main asymmetry (beta-adjusted returns, a 3-day window, odds read 24 hours and
    7 days earlier, with Welch t-statistics) and a tougher benchmark: a logistic regression on each company's analyst
    history, trained only on earlier months. Reads everything it needs from the workbook."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as L
    from openpyxl.worksheet.formula import ArrayFormula
    BOLD = Font(name="Arial", bold=True); NOTE = Font(name="Arial", italic=True, color="666666")
    SEC = Font(name="Arial", bold=True, size=11, color="1F3A5F"); HEAD = PatternFill("solid", fgColor="DDE4EE")

    def head(ws, r, vals, c0=1):
        for j, v in enumerate(vals):
            c = ws.cell(row=r, column=c0 + j, value=v); c.font = BOLD; c.fill = HEAD
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 45

    EVT, RES, RP, RM, RE = (wb[n] for n in ("Events", "Results", "Raw_Prices", "Raw_Markets", "Raw_Estimates"))
    PE = 4
    while RP.cell(row=PE + 1, column=1).value is not None:
        PE += 1
    pcol = {RP.cell(row=4, column=j).value: L(j) for j in range(2, RP.max_column + 1)}
    E_ = 4
    while EVT.cell(row=E_ + 1, column=1).value is not None:
        E_ += 1
    cM = pcol.get(MARKET)
    if E_ < 6 or not cM:
        return
    head(EVT, 4, [f"Beta ({beta_days} days, ending {gap_days} days before)", "Beta-adjusted abnormal return", "Abnormal return, 3-day window",
                  "p 24 hours before the cut-off", "Analyst-history model: P(beat)", "Brier: analyst-history model",
                  "Helper: abnormal squared", "Helper: beta-adjusted squared", "Helper: 3-day squared"], c0=53)
    M_ = f"Raw_Prices!${cM}$5:${cM}${PE}"
    for r in range(5, E_ + 1):
        raw = EVT.cell(row=r, column=1).value
        cS = pcol.get(yf_symbol(str(RM.cell(row=raw, column=3).value)))
        a, b = EVT.cell(row=r, column=6).value, EVT.cell(row=r, column=7).value
        if cS:
            S_ = f"Raw_Prices!${cS}$5:${cS}${PE}"
            pre = f"MATCH(C{r}-1,Raw_Prices!$A$5:$A${PE},1)"
            rng = lambda X, k0, k1: f"INDEX({X},{pre}-{k0}):INDEX({X},{pre}-{k1})"
            hi_, lo_ = gap_days + beta_days - 1, gap_days
            EVT[f"BA{r}"] = ArrayFormula(f"BA{r}", (f'=IFERROR(SLOPE(LN({rng(S_, hi_, lo_)}/{rng(S_, hi_ + 1, lo_ + 1)}),'
                                                    f'LN({rng(M_, hi_, lo_)}/{rng(M_, hi_ + 1, lo_ + 1)})),"")'))
            EVT[f"BA{r}"].number_format = "0.00"
            post2 = f"(MATCH(C{r},Raw_Prices!$A$5:$A${PE},1)+2)"
            EVT[f"BC{r}"] = (f'=IFERROR(LN(INDEX({S_},{post2})/INDEX({S_},{pre}))-LN(INDEX({M_},{post2})/INDEX({M_},{pre})),"")')
        EVT[f"BB{r}"] = f'=IF(AND(ISNUMBER(S{r}),ISNUMBER(T{r}),ISNUMBER(BA{r})),S{r}-BA{r}*T{r},"")'
        if a and b:
            EVT[f"BD{r}"] = f'=IFERROR(INDEX(Raw_Odds!$D${a}:$D${b},MATCH(H{r}-86400,Raw_Odds!$B${a}:$B${b},1)),"")'
        EVT[f"BF{r}"] = f'=IF(ISNUMBER(BE{r}),(BE{r}-E{r})^2,"")'
        EVT[f"BG{r}"] = f'=IF(ISNUMBER(U{r}),U{r}^2,"")'
        EVT[f"BH{r}"] = f'=IF(ISNUMBER(BB{r}),BB{r}^2,"")'
        EVT[f"BI{r}"] = f'=IF(ISNUMBER(BC{r}),BC{r}^2,"")'
        for c_ in ("BB", "BC"):
            EVT[f"{c_}{r}"].number_format = "+0.00%;-0.00%"
        for c_ in ("BD", "BE", "BF"):
            EVT[f"{c_}{r}"].number_format = "0.000"

    # analyst-history model, trained month by month on earlier events only
    try:
        from sklearn.linear_model import LogisticRegression
        EE = 4
        while RE.cell(row=EE + 1, column=1).value is not None:
            EE += 1
        est = pd.DataFrame([[RE.cell(row=i, column=c).value for c in range(1, 6)] for i in range(5, EE + 1)],
                           columns=["t", "d", "est", "rep", "sp"])
        est["d"] = pd.to_datetime(est["d"]); est["b"] = (pd.to_numeric(est["rep"]) > pd.to_numeric(est["est"])).astype(int)
        est["sp"] = pd.to_numeric(est["sp"], errors="coerce").clip(-100, 100)
        ev = []
        for r in range(5, E_ + 1):
            raw = EVT.cell(row=r, column=1).value
            t, d, y = str(RM.cell(row=raw, column=3).value), pd.Timestamp(RM.cell(row=raw, column=4).value), RM.cell(row=raw, column=9).value
            h = est[(est["t"] == t) & (est["d"] < d - pd.Timedelta(days=3))].sort_values("d").tail(12)
            if len(h) >= MIN_HISTORY and y in (0, 1):
                streak = int(h["b"][::-1].cumprod().sum())
                ev.append(dict(r=r, d=d, y=int(y), br=h["b"].mean(), ms=float(h["sp"].median()) if h["sp"].notna().any() else 0.0, streak=streak))
        df = pd.DataFrame(ev)
        if len(df):
            lg = lambda x: np.log(np.clip(x, .02, .98) / (1 - np.clip(x, .02, .98)))
            df["lbr"] = lg(df["br"]); df["m"] = df["d"].dt.to_period("M")
            F = ["lbr", "ms", "streak"]
            for m in sorted(df["m"].unique()):
                tr, te = df[df["m"] < m], df[df["m"] == m]
                if len(tr) < min_train or tr["y"].nunique() < 2:
                    continue
                mdl = LogisticRegression(max_iter=1000).fit(tr[F], tr["y"])
                for rr_, pr in zip(te["r"], mdl.predict_proba(te[F])[:, 1]):
                    EVT[f"BE{rr_}"] = round(float(pr), 4)
    except Exception as e:
        print(f"   analyst-history model skipped: {e}", flush=True)

    R = lambda c: f"Events!${c}$5:${c}${E_}"
    P2 = "+0.00%;-0.00%"
    rr = RES.max_row + 3
    RES.cell(row=rr, column=1, value="Q. Robustness of the main result: do confident misses still fall more?").font = SEC; rr += 1
    head(RES, rr, ["Measure", "Miss, p at least 0.8", "Events", "Miss, p below 0.8", "Events", "Difference", "Welch t",
                   "Beat, p at least 0.8", "Beat, p below 0.8"]); rr += 1
    for lab, ret, sq, pc in (("Market-adjusted return, odds at the cut-off (main result)", "U", "BG", "I"),
                             (f"Beta-adjusted return ({beta_days}-day beta)", "BB", "BH", "I"),
                             ("Market-adjusted return, 3-day window", "BC", "BI", "I"),
                             ("Odds read 24 hours earlier", "U", "BG", "BD"),
                             (f"Odds read {LEAD_DAYS} days earlier", "U", "BG", "J")):
        RES.cell(row=rr, column=1, value=lab)
        for j, op in enumerate(('">=0.8"', '"<0.8"')):
            cr = f'{R("E")},0,{R(pc)},{op}'
            RES.cell(row=rr, column=2 + 2 * j, value=f'=IFERROR(AVERAGEIFS({R(ret)},{cr}),"n/a")').number_format = P2
            RES.cell(row=rr, column=3 + 2 * j, value=f'=COUNTIFS({cr},{R(ret)},">=-10")')
        RES.cell(row=rr, column=6, value=f'=IFERROR(B{rr}-D{rr},"n/a")').number_format = P2
        v = lambda mc, nc, op: f'(SUMIFS({R(sq)},{R("E")},0,{R(pc)},{op})-{nc}{rr}*{mc}{rr}^2)/({nc}{rr}-1)'
        RES.cell(row=rr, column=7, value=(f'=IFERROR(F{rr}/SQRT({v("B", "C", chr(34) + ">=0.8" + chr(34))}/C{rr}'
                                          f'+{v("D", "E", chr(34) + "<0.8" + chr(34))}/E{rr}),"n/a")')).number_format = "0.00"
        for j, op in enumerate(('">=0.8"', '"<0.8"')):
            RES.cell(row=rr, column=8 + j, value=f'=IFERROR(AVERAGEIFS({R(ret)},{R("E")},1,{R(pc)},{op}),"n/a")').number_format = P2
        rr += 1
    for k, (lab, lo, hi) in enumerate((("Lowest third by volatility before the report", 0, 1 / 3), ("Middle third", 1 / 3, 2 / 3),
                                       ("Highest third", 2 / 3, 1))):
        vc = f'{R("AM")},">="&PERCENTILE({R("AM")},{lo}),{R("AM")},"{"<=" if hi == 1 else "<"}"&PERCENTILE({R("AM")},{hi})'
        RES.cell(row=rr, column=1, value=lab)
        for j, op in enumerate(('">=0.8"', '"<0.8"')):
            cr = f'{R("E")},0,{R("I")},{op},{vc}'
            RES.cell(row=rr, column=2 + 2 * j, value=f'=IFERROR(AVERAGEIFS({R("U")},{cr}),"n/a")').number_format = P2
            RES.cell(row=rr, column=3 + 2 * j, value=f'=COUNTIFS({cr},{R("U")},">=-10")')
        RES.cell(row=rr, column=6, value=f'=IFERROR(B{rr}-D{rr},"n/a")').number_format = P2
        rr += 1
    RES.cell(row=rr, column=1, value=("Reading: the main result should not depend on the return measure or the window. With odds read earlier "
                                      "the gap is expected to be smaller, because the crowd keeps learning until the release. Event counts here include only "
                                      "reports with price data for the window, so they can be lower than in part C.")).font = NOTE
    rr += 2
    RES.cell(row=rr, column=1, value="Tougher benchmark: logistic regression on each company's analyst history (beat rate, median surprise, "
                                     "beat streak over its last 12 reports), trained only on earlier months").font = SEC; rr += 1
    head(RES, rr, ["", "Value"]); rr += 1
    q0 = rr
    for lab, fm, nf in (("Events with a model forecast", f"=COUNT({R('BE')})", "0"),
                        ("Brier, crowd (same events)", f'=IFERROR(SUMIFS({R("O")},{R("BE")},">=0")/COUNTIFS({R("BE")},">=0",{R("O")},">=0"),"n/a")', "0.000"),
                        ("Brier, analyst-history model", f'=IFERROR(AVERAGE({R("BF")}),"n/a")', "0.000"),
                        ("Skill of the crowd vs the model", f'=IFERROR(1-B{q0 + 1}/B{q0 + 2},"n/a")', "0.0%")):
        RES.cell(row=rr, column=1, value=lab); RES.cell(row=rr, column=2, value=fm).number_format = nf; rr += 1
    RES.cell(row=rr, column=1, value=("The model's forecasts are written by the script (scikit-learn); the comparison is done with formulas. "
                                      "Positive skill means the crowd beats the best forecast we could build from analysts' track records.")).font = NOTE


def add_lookup(wb, max_rows=20):
    """Company lookup sheet: type a ticker and see that company's past reports, its record and, if it reports soon,
    its scenario card from the Live sheet. Plain formulas (no array formulas), so it works in any version of Excel."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    BOLD = Font(name="Arial", bold=True); NOTE = Font(name="Arial", italic=True, color="666666")
    TITLE = Font(name="Arial", bold=True, size=14, color="1F3A5F"); SEC = Font(name="Arial", bold=True, size=11, color="1F3A5F")
    HEAD = PatternFill("solid", fgColor="DDE4EE"); INPUT = PatternFill("solid", fgColor="FFF2CC")
    EVT, LIV, RM = wb["Events"], wb["Live"], wb["Raw_Markets"]
    E_ = 4
    while EVT.cell(row=E_ + 1, column=1).value is not None:
        E_ += 1
    if E_ < 6:
        return
    tick = sorted({str(RM.cell(row=EVT.cell(row=r, column=1).value, column=3).value) for r in range(5, E_ + 1)})
    LK = wb.create_sheet("Company", 1)
    LK["A1"] = "Company lookup: one company's earnings history and scenario"; LK["A1"].font = TITLE
    LK["A3"] = "Type or pick a ticker:"; LK["A3"].font = BOLD
    LK["B3"] = "JPM" if "JPM" in tick else tick[0]; LK["B3"].fill = INPUT; LK["B3"].font = Font(name="Arial", bold=True, color="0000FF")
    LK["C3"] = "(yellow cell; the drop-down lists every company in the sample)"; LK["C3"].font = NOTE
    for i, t in enumerate(tick):
        LK.cell(row=5 + i, column=26, value=t)                       # list for the drop-down, column Z
    LK["Z4"] = "Tickers"; LK.column_dimensions["Z"].hidden = True
    dv = DataValidation(type="list", formula1=f"=$Z$5:$Z${4 + len(tick)}", allow_blank=False)
    LK.add_data_validation(dv); dv.add("B3")
    # helper on Events: running count of the selected company's reports
    EVT.cell(row=4, column=63, value="Helper: n-th report of the company selected on the Company sheet").font = BOLD
    for r in range(5, E_ + 1):
        EVT[f"BK{r}"] = f'=IF(B{r}=Company!$B$3,COUNTIF(B$5:B{r},Company!$B$3),"")'
    R = lambda c: f"Events!${c}$5:${c}${E_}"
    T = "$B$3"
    rows = [("Past reports in the sample", f"=COUNTIF({R('B')},{T})", "0"),
            ("Share that beat", f'=IFERROR(AVERAGEIFS({R("E")},{R("B")},{T}),"n/a")', "0%"),
            ("Average crowd probability before the report", f'=IFERROR(AVERAGEIFS({R("I")},{R("B")},{T}),"n/a")', "0%"),
            ("Company's beat record (Yahoo, past reports)", f'=IFERROR(AVERAGEIFS({R("M")},{R("B")},{T}),"n/a")', "0%"),
            ("Average move on the report (|abnormal return|)", f'=IFERROR(AVERAGEIFS({R("X")},{R("B")},{T}),"n/a")', "0.0%"),
            ("Average move when it beat", f'=IFERROR(AVERAGEIFS({R("U")},{R("B")},{T},{R("E")},1),"n/a")', "+0.0%;-0.0%"),
            ("Average move when it missed", f'=IFERROR(AVERAGEIFS({R("U")},{R("B")},{T},{R("E")},0),"n/a")', "+0.0%;-0.0%"),
            ("Misses when the crowd was at least 80% sure of a beat", f'=COUNTIFS({R("B")},{T},{R("E")},0,{R("I")},">=0.8")', "0"),
            ("Crowd Brier score on this company (lower is better)", f'=IFERROR(AVERAGEIFS({R("O")},{R("B")},{T}),"n/a")', "0.000")]
    LK["A5"] = "This company's record"; LK["A5"].font = SEC
    for k, (lab, fm, nf) in enumerate(rows):
        LK.cell(row=6 + k, column=1, value=lab); LK.cell(row=6 + k, column=2, value=fm).number_format = nf
    # upcoming report from the Live sheet
    n_live = 0
    while LIV.cell(row=5 + n_live, column=1).value is not None:
        n_live += 1
    r0 = 17
    LK.cell(row=r0, column=1, value="Upcoming report (from the Live sheet)").font = SEC
    if n_live:
        L1 = 4 + n_live
        m = f"MATCH({T},Live!$A$5:$A${L1},0)"
        live_rows = [("Report date", "B", "yyyy-mm-dd"), ("Crowd probability of a beat now", "F", "0%"),
                     ("Gap from the company's record", "H", "+0%;-0%"), ("Expected size of the move", "R", "0.0%"),
                     ("Past move if it beats (all companies, same crowd confidence)", "S", "+0.0%;-0.0%"),
                     ("Past move if it misses (all companies, same crowd confidence)", "T", "+0.0%;-0.0%"),
                     ("Option-implied move", "K", "0.0%"), ("Reading", "Y", "@")]
        for k, (lab, col, nf) in enumerate(live_rows):
            LK.cell(row=r0 + 1 + k, column=1, value=lab)
            LK.cell(row=r0 + 1 + k, column=2, value=f'=IFERROR(INDEX(Live!${col}$5:${col}${L1},{m}),"no upcoming report listed")').number_format = nf
    else:
        LK.cell(row=r0 + 1, column=1, value="No upcoming reports in this run.")
    # past reports list
    t0 = 28
    LK.cell(row=t0, column=1, value=f"Past reports (up to {max_rows}, oldest first)").font = SEC
    hdr = ["Report no.", "Report date", "Crowd p before", "Beat (1) or miss (0)", "Abnormal return", "Analyst surprise (%)", "Company record at the time"]
    for j, h in enumerate(hdr):
        c = LK.cell(row=t0 + 1, column=1 + j, value=h); c.font = BOLD; c.fill = HEAD; c.alignment = Alignment(wrap_text=True, vertical="top")
    LK.row_dimensions[t0 + 1].height = 30
    for k in range(1, max_rows + 1):
        r = t0 + 1 + k
        LK.cell(row=r, column=1, value=f'=IF({k}<=$B$6,{k},"")')
        m = f"MATCH({k},{R('BK')},0)"
        for j, (col, nf) in enumerate((("C", "yyyy-mm-dd"), ("I", "0%"), ("E", "0"), ("U", "+0.0%;-0.0%"), ("N", "0.0"), ("M", "0%"))):
            LK.cell(row=r, column=2 + j, value=f'=IFERROR(INDEX({R(col)},{m}),"")').number_format = nf
    LK.cell(row=t0 + max_rows + 3, column=1, value=(
        "Reading: a single company has only a few reports in 13 months, so its own averages are noisy. The 'same crowd "
        "confidence' rows use all companies, which is where the evidence is.")).font = NOTE
    LK.column_dimensions["A"].width = 58
    for c in "BCDEFG":
        LK.column_dimensions[c].width = 16
    LK.sheet_properties.tabColor = "2E7D32"


def add_cutoffs(wb, cutoffs=(0.70, 0.75, 0.80, 0.85, 0.90)):
    """Results part R: how strict should 'confident' be? The confident-miss gap at several cut-offs, with a 95% interval
    (Welch) and p-value. All formulas. Needs the squared-return helper written by add_robust (Events column BG)."""
    from openpyxl.styles import Alignment, Font, PatternFill
    BOLD = Font(name="Arial", bold=True); NOTE = Font(name="Arial", italic=True, color="666666")
    SEC = Font(name="Arial", bold=True, size=11, color="1F3A5F"); HEAD = PatternFill("solid", fgColor="DDE4EE")
    BLUE = Font(name="Arial", color="0000FF")
    EVT, RES = wb["Events"], wb["Results"]
    E_ = 4
    while EVT.cell(row=E_ + 1, column=1).value is not None:
        E_ += 1
    if E_ < 6:
        return
    R = lambda c: f"Events!${c}$5:${c}${E_}"
    rr = RES.max_row + 3
    RES.cell(row=rr, column=1, value="R. How strict should 'confident' be? Misses at different cut-offs of the crowd probability").font = SEC; rr += 1
    hdr = ["Cut-off (crowd p at least)", "Confident misses", "Their average move", "Other misses: average move", "Difference",
           "95% interval: low", "95% interval: high", "Welch t", "p-value", "Helper: variance, confident", "Helper: variance, other",
           "Helper: degrees of freedom"]
    for j, h in enumerate(hdr):
        c = RES.cell(row=rr, column=1 + j, value=h); c.font = BOLD; c.fill = HEAD; c.alignment = Alignment(wrap_text=True, vertical="top")
    RES.row_dimensions[rr].height = 45; rr += 1
    P2 = "+0.00%;-0.00%"
    for cut in cutoffs:
        A = f"$A{rr}"
        hi_, lo_ = f'">="&{A}', f'"<"&{A}'
        RES.cell(row=rr, column=1, value=cut).font = BLUE
        RES.cell(row=rr, column=1).number_format = "0.00"
        RES.cell(row=rr, column=2, value=f'=COUNTIFS({R("E")},0,{R("I")},{hi_},{R("U")},">=-10")')
        RES.cell(row=rr, column=3, value=f'=IFERROR(AVERAGEIFS({R("U")},{R("E")},0,{R("I")},{hi_}),"n/a")').number_format = P2
        RES.cell(row=rr, column=4, value=f'=IFERROR(AVERAGEIFS({R("U")},{R("E")},0,{R("I")},{lo_}),"n/a")').number_format = P2
        RES.cell(row=rr, column=5, value=f'=IFERROR(C{rr}-D{rr},"n/a")').number_format = P2
        n_lo = f'COUNTIFS({R("E")},0,{R("I")},{lo_},{R("U")},">=-10")'
        RES.cell(row=rr, column=10, value=f'=IFERROR((SUMIFS({R("BG")},{R("E")},0,{R("I")},{hi_})-B{rr}*C{rr}^2)/(B{rr}-1),"n/a")')
        RES.cell(row=rr, column=11, value=f'=IFERROR((SUMIFS({R("BG")},{R("E")},0,{R("I")},{lo_})-{n_lo}*D{rr}^2)/({n_lo}-1),"n/a")')
        se = f"SQRT(J{rr}/B{rr}+K{rr}/{n_lo})"
        RES.cell(row=rr, column=12, value=(f'=IFERROR((J{rr}/B{rr}+K{rr}/{n_lo})^2/((J{rr}/B{rr})^2/(B{rr}-1)'
                                           f'+(K{rr}/{n_lo})^2/({n_lo}-1)),"n/a")')).number_format = "0.0"
        RES.cell(row=rr, column=6, value=f'=IFERROR(E{rr}-_xlfn.T.INV.2T(0.05,L{rr})*{se},"n/a")').number_format = P2
        RES.cell(row=rr, column=7, value=f'=IFERROR(E{rr}+_xlfn.T.INV.2T(0.05,L{rr})*{se},"n/a")').number_format = P2
        RES.cell(row=rr, column=8, value=f'=IFERROR(E{rr}/{se},"n/a")').number_format = "0.00"
        RES.cell(row=rr, column=9, value=f'=IFERROR(_xlfn.T.DIST.2T(ABS(H{rr}),L{rr}),"n/a")').number_format = "0.000"
        for c_ in (10, 11):
            RES.cell(row=rr, column=c_).number_format = "0.00000"
        rr += 1
    RES.cell(row=rr, column=1, value=("Reading: 0.80 was fixed in the first version of the tool, before any results. A stricter cut-off leaves "
                                      "fewer misses, so the interval widens. Cut-offs are blue inputs and can be changed.")).font = NOTE


def build_fed(wb, fed, download_utc, head, st):
    """Fed sheets: how well the crowd priced each Fed decision, and whether the earnings results differ around Fed weeks
    or when 2-year yields moved. All statistics are formulas."""
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter as L
    from openpyxl.worksheet.formula import ArrayFormula
    BOLD, NOTE, TITLE, SEC, BLUE = st["BOLD"], st["NOTE"], st["TITLE"], st["SEC"], st["BLUE"]
    fm = fed["markets"].copy()
    fm = fm.sort_values(["date", "venue", "outcome"]).reset_index(drop=True)
    i0 = wb.sheetnames.index("Assumptions")
    FR, FE = wb.create_sheet("Fed", i0), wb.create_sheet("Fed_Events", i0 + 1)
    RFM, RFO = wb.create_sheet("Raw_Fed_Markets"), wb.create_sheet("Raw_Fed_Odds")
    RY = wb.create_sheet("Raw_Rates")
    # ---- raw ----
    RFM["A1"] = (f"Fed decision contracts as downloaded: Kalshi {KALSHI} (series {', '.join(FED_SERIES)}) and Polymarket {GAMMA}/events "
                 f"(Fed decision events). Downloaded {download_utc} UTC."); RFM["A1"].font = BOLD
    head(RFM, 4, ["Venue", "Meeting", "Decision date", "Outcome", "Result (1 = happened)", "Volume", "Close (unix)", "Contract", "Price now (open contracts)"])
    for i, x in enumerate(fm.itertuples()):
        r = 5 + i
        vals = [x.venue, x.meeting, x.date.to_pydatetime(), x.outcome, (None if pd.isna(x.y) else int(x.y)), float(x.volume), int(x.close_ts), x.market,
                (x.p_now if (x.open_ and x.p_now is not None and not pd.isna(x.p_now)) else None)]
        for j, v in enumerate(vals):
            RFM.cell(row=r, column=1 + j, value=v)
        RFM.cell(row=r, column=3).number_format = "yyyy-mm-dd"
    RFO["A1"] = "Hourly YES prices for each resolved Fed contract, 32 days before the decision (Kalshi candlesticks, Polymarket CLOB)."; RFO["A1"].font = BOLD
    head(RFO, 4, ["Contract row (Raw_Fed_Markets)", "Time (unix)", "Time (UTC)", "YES price"])
    odict = dict(fed["odds"])
    r, span = 5, {}
    for i, x in enumerate(fm.itertuples()):
        a = r
        for t, pz in odict.get(x.market, []):
            RFO.cell(row=r, column=1, value=5 + i); RFO.cell(row=r, column=2, value=t)
            RFO.cell(row=r, column=3, value=datetime.fromtimestamp(t, timezone.utc).replace(tzinfo=None)).number_format = "yyyy-mm-dd hh:mm"
            RFO.cell(row=r, column=4, value=pz); r += 1
        span[i] = (a, r - 1) if r > a else None
    if "Yahoo" in fed.get("yname", ""):
        RY["A1"] = (f"Yahoo Finance via yfinance, ticker ^FVX: 5-year Treasury yield (%), daily. FRED ({FRED_CSV}) was not reachable "
                    f"in this run, so the Fed target column is empty. Downloaded {download_utc} UTC.")
    else:
        RY["A1"] = (f"FRED public CSV ({FRED_CSV}): DGS2 = 2-year Treasury yield (%), DFEDTARU = Fed funds target, upper bound (%). "
                    f"Downloaded {download_utc} UTC.")
    RY["A1"].font = BOLD
    ASM = wb["Assumptions"]
    ar = ASM.max_row + 2
    ASM.cell(row=ar, column=1, value="Fed decisions and interest rates").font = SEC; ar += 1
    for lab, v in [("Kalshi Fed contracts", f"{KALSHI}: series {', '.join(FED_SERIES)}, hourly candlesticks. Downloaded {download_utc} UTC"),
                   ("Polymarket Fed contracts", f"{GAMMA}/events (tags {', '.join(FED_TAGS)}) and {CLOB}/prices-history, hourly. Downloaded {download_utc} UTC"),
                   ("Interest rate", f"{fed.get('yname', '')}, daily. Downloaded {download_utc} UTC"),
                   ("Fed meetings from", FED_START), ("Fed week: days from a decision", FED_WEEK_DAYS),
                   ("Yield change counted as rising or falling (points, 20 trading days)", YIELD_MOVE)]:
        ASM.cell(row=ar, column=1, value=lab); c = ASM.cell(row=ar, column=2, value=v)
        if isinstance(v, (int, float)):
            c.font = BLUE
        ar += 1
    head(RY, 4, ["Date", fed.get("yname", "2-year yield"), "", "Date", "Fed target upper (DFEDTARU)"])
    y2, tg = fed["y2"].reset_index(drop=True), fed["target"].reset_index(drop=True)
    for i, x in enumerate(y2.itertuples()):
        RY.cell(row=5 + i, column=1, value=x.date.to_pydatetime()).number_format = "yyyy-mm-dd"; RY.cell(row=5 + i, column=2, value=float(x.value))
    for i, x in enumerate(tg.itertuples()):
        RY.cell(row=5 + i, column=4, value=x.date.to_pydatetime()).number_format = "yyyy-mm-dd"; RY.cell(row=5 + i, column=5, value=float(x.value))
    YE, TE = 4 + max(1, len(y2)), 4 + max(1, len(tg))
    for ws in (RFM, RFO, RY):
        ws.freeze_panes = "A5"; ws.sheet_properties.tabColor = "C65911"

    # ---- Fed_Events: one row per resolved contract with odds, grouped by meeting and venue ----
    FE["A1"] = "Fed_Events: one row per resolved Fed outcome contract (formulas linking to Raw_Fed_ sheets)"; FE["A1"].font = TITLE
    FE["A2"] = ("p = YES price 30, 7 and 1 days before the decision. Squared error = (p - result)^2, summed over the outcomes of a meeting "
                "gives the meeting's Brier score. Favourite = the outcome with the highest p in its meeting and venue. "
                "'No change' benchmark = always predict a hold."); FE["A2"].alignment = Alignment(wrap_text=True)
    FE.merge_cells("A2:U2"); FE.row_dimensions[2].height = 40
    head(FE, 4, ["Raw row", "Venue", "Meeting", "Decision date", "Outcome", "Result y", "Odds first row", "Odds last row", "Close (unix)",
                 "p 30 days before", "p 7 days before", "p 1 day before", "Squared error, 1 day", "Squared error, 7 days", "Squared error, 30 days",
                 "Favourite 1 day before (1/0)", "Favourite 7 days before (1/0)", "Group first row", "Group last row",
                 "Squared error, 'no change' benchmark", "Meeting and venue key"])
    good = [(i, span[i]) for i, x in enumerate(fm.itertuples()) if span.get(i) and not pd.isna(x.y)]
    keys = {}
    for k, (i, _) in enumerate(good):
        x = fm.iloc[i]; keys.setdefault((x["meeting"], x["venue"]), []).append(5 + k)
    for k, (i, (a, b)) in enumerate(good):
        r = 5 + k; raw = 5 + i
        x = fm.iloc[i]; grp = keys[(x["meeting"], x["venue"])]; g0, g1 = min(grp), max(grp)
        f = {"A": raw, "B": f"=Raw_Fed_Markets!A{raw}", "C": f"=Raw_Fed_Markets!B{raw}", "D": f"=Raw_Fed_Markets!C{raw}",
             "E": f"=Raw_Fed_Markets!D{raw}", "F": f"=Raw_Fed_Markets!E{raw}", "G": a, "H": b, "I": f"=Raw_Fed_Markets!G{raw}"}
        for col, days in (("J", 30), ("K", 7), ("L", 1)):
            f[col] = f'=IFERROR(INDEX(Raw_Fed_Odds!$D${a}:$D${b},MATCH(I{r}-{days}*86400,Raw_Fed_Odds!$B${a}:$B${b},1)),"")'
        f["M"] = f'=IF(ISNUMBER(L{r}),(L{r}-F{r})^2,"")'; f["N"] = f'=IF(ISNUMBER(K{r}),(K{r}-F{r})^2,"")'; f["O"] = f'=IF(ISNUMBER(J{r}),(J{r}-F{r})^2,"")'
        f["P"] = f'=IF(ISNUMBER(L{r}),IF(L{r}=MAX(L${g0}:L${g1}),1,0),"")'
        f["Q"] = f'=IF(ISNUMBER(K{r}),IF(K{r}=MAX(K${g0}:K${g1}),1,0),"")'
        f["R"] = g0; f["S"] = g1
        f["T"] = f'=(IF(E{r}="hold",1,0)-F{r})^2'
        f["U"] = f'=C{r}&"|"&B{r}'
        for c_, v in f.items():
            FE[f"{c_}{r}"] = v
        FE[f"D{r}"].number_format = "yyyy-mm-dd"
        for c_ in "JKLMNOT":
            FE[f"{c_}{r}"].number_format = "0.000"
    E2 = 4 + max(1, len(good)); FE.freeze_panes = "C5"
    G = lambda c: f"Fed_Events!${c}$5:${c}${E2}"

    # ---- Fed results ----
    FR["A1"] = "Fed meetings: how well did the crowd price each decision?"; FR["A1"].font = TITLE
    FR["A2"] = ("Each Fed meeting has five outcome contracts (hold, cut 25bp, cut more than 25bp, hike 25bp, hike more than 25bp). "
                "We read the crowd's probabilities 30, 7 and 1 days before the decision and compare them with what the Fed did. "
                "Brier per meeting = sum over outcomes of (p - result)^2, lower is better. The benchmark always predicts 'no change'.")
    FR["A2"].alignment = Alignment(wrap_text=True); FR.merge_cells("A2:G2"); FR.row_dimensions[2].height = 55
    head(FR, 4, ["", "Kalshi", "Polymarket"])
    labels = ["Meetings", "Average p on what happened, 30 days before", "Average p on what happened, 7 days before",
              "Average p on what happened, 1 day before", "Favourite won, 7 days before", "Favourite won, 1 day before",
              "Brier per meeting, 1 day before", "Brier per meeting, 7 days before", "Brier per meeting, 'no change' benchmark",
              "Skill vs 'no change', 7 days before", "Meetings where the outcome had p below 0.5 one day before (surprises)"]
    for k, lab in enumerate(labels):
        FR.cell(row=5 + k, column=1, value=lab)
    for j, v in enumerate(("Kalshi", "Polymarket")):
        c = L(2 + j); vv = f'"{v}"'
        fms = [f'=COUNTIFS({G("B")},{vv},{G("F")},1,{G("L")},">=0")',
               f'=IFERROR(AVERAGEIFS({G("J")},{G("B")},{vv},{G("F")},1),"n/a")',
               f'=IFERROR(AVERAGEIFS({G("K")},{G("B")},{vv},{G("F")},1),"n/a")',
               f'=IFERROR(AVERAGEIFS({G("L")},{G("B")},{vv},{G("F")},1),"n/a")',
               f'=IFERROR(COUNTIFS({G("B")},{vv},{G("Q")},1,{G("F")},1)/{c}5,"n/a")',
               f'=IFERROR(COUNTIFS({G("B")},{vv},{G("P")},1,{G("F")},1)/{c}5,"n/a")',
               f'=IFERROR(SUMIFS({G("M")},{G("B")},{vv})/{c}5,"n/a")',
               f'=IFERROR(SUMIFS({G("N")},{G("B")},{vv})/{c}5,"n/a")',
               f'=IFERROR(SUMIFS({G("T")},{G("B")},{vv},{G("L")},">=0")/{c}5,"n/a")',
               f'=IFERROR(1-{c}12/{c}13,"n/a")',
               f'=COUNTIFS({G("B")},{vv},{G("F")},1,{G("L")},"<0.5")']
        nfs = ["0", "0.0%", "0.0%", "0.0%", "0.0%", "0.0%", "0.000", "0.000", "0.000", "0.0%", "0"]
        for k, (f_, nf) in enumerate(zip(fms, nfs)):
            FR.cell(row=5 + k, column=2 + j, value=f_).number_format = nf
    # per-meeting table
    rr = 18
    FR.cell(row=rr, column=1, value="Meeting by meeting: probability the crowd gave to what actually happened").font = SEC; rr += 1
    head(FR, rr, ["Decision date", "What the Fed did", "Kalshi, 30 days before", "Kalshi, 7 days before", "Kalshi, 1 day before",
                  "Polymarket, 7 days before", "Polymarket, 1 day before"]); rr += 1
    yes_rows = {}
    for k, (i, _) in enumerate(good):
        x = fm.iloc[i]
        if x["y"] == 1:
            yes_rows.setdefault(x["date"], {})[x["venue"]] = 5 + k
    m0 = rr
    for d in sorted(yes_rows):
        vr = yes_rows[d]; any_r = vr.get("Kalshi") or vr.get("Polymarket")
        FR.cell(row=rr, column=1, value=f"=Fed_Events!D{any_r}").number_format = "yyyy-mm-dd"
        FR.cell(row=rr, column=2, value=f"=Fed_Events!E{any_r}")
        for col, (v, src) in zip(range(3, 8), (("Kalshi", "J"), ("Kalshi", "K"), ("Kalshi", "L"), ("Polymarket", "K"), ("Polymarket", "L"))):
            if v in vr:
                FR.cell(row=rr, column=col, value=f"=Fed_Events!{src}{vr[v]}").number_format = "0%"
        rr += 1
    m1 = rr - 1
    for c_, w in zip("ABCDEFG", (58, 18, 14, 14, 14, 14, 14)):
        FR.column_dimensions[c_].width = w

    # ---- link to the earnings results: Events columns and Results part P ----
    EVT, RES = wb["Events"], wb["Results"]
    E_ = 4
    while EVT.cell(row=E_ + 1, column=1).value is not None:
        E_ += 1
    if E_ < 6 or m1 < m0:
        return
    head(EVT, 4, ["Days to the nearest Fed decision", f"Change in {fed.get('yname', '2-year yield')}, 20 trading days before (points)", "Fed target at the report (%)"], c0=49)
    for r in range(5, E_ + 1):
        EVT[f"AW{r}"] = ArrayFormula(f"AW{r}", f"=MIN(ABS(Fed!$A${m0}:$A${m1}-C{r}))")
        if len(y2):
            mt = f"MATCH(C{r}-1,Raw_Rates!$A$5:$A${YE},1)"
            EVT[f"AX{r}"] = f'=IFERROR(INDEX(Raw_Rates!$B$5:$B${YE},{mt})-INDEX(Raw_Rates!$B$5:$B${YE},{mt}-20),"")'
            EVT[f"AX{r}"].number_format = "+0.00;-0.00"
        if len(tg):
            EVT[f"AY{r}"] = f'=IFERROR(INDEX(Raw_Rates!$E$5:$E${TE},MATCH(C{r},Raw_Rates!$D$5:$D${TE},1)),"")'
    R = lambda c: f"Events!${c}$5:${c}${E_}"
    P2 = "+0.00%;-0.00%"
    rr = RES.max_row + 3
    RES.cell(row=rr, column=1, value="P. Rates and the Fed: do the earnings results change around Fed decisions or when yields move?").font = SEC; rr += 1
    head(RES, rr, ["", "Events", "Skill vs base rate", "Average abnormal return: miss, p at least 0.8", "Misses, p at least 0.8",
                   "Average abnormal return: beat, p at least 0.8", "Average |abnormal return|"]); rr += 1
    groups = (("All events", None, ""),
              (f"Within {FED_WEEK_DAYS} days of a Fed decision", f"--({R('AW')}<={FED_WEEK_DAYS})", f',{R("AW")},"<={FED_WEEK_DAYS}"'),
              ("Other weeks", f"--({R('AW')}>{FED_WEEK_DAYS})", f',{R("AW")},">{FED_WEEK_DAYS}"'),
              (f"Treasury yield rose by {YIELD_MOVE} or more", f"ISNUMBER({R('AX')})*({R('AX')}>={YIELD_MOVE})", f',{R("AX")},">={YIELD_MOVE}"'),
              ("Treasury yield roughly flat", f"ISNUMBER({R('AX')})*({R('AX')}>-{YIELD_MOVE})*({R('AX')}<{YIELD_MOVE})",
               f',{R("AX")},">-{YIELD_MOVE}",{R("AX")},"<{YIELD_MOVE}"'),
              (f"Treasury yield fell by {YIELD_MOVE} or more", f"ISNUMBER({R('AX')})*({R('AX')}<=-{YIELD_MOVE})", f',{R("AX")},"<=-{YIELD_MOVE}"'))
    for lab, m, crit in groups:
        mm = f",{m}" if m else ""
        RES.cell(row=rr, column=1, value=lab)
        RES.cell(row=rr, column=2, value=f"=SUMPRODUCT(--ISNUMBER({R('E')}){mm})")
        RES.cell(row=rr, column=3, value=f'=IFERROR(1-SUMPRODUCT({R("R")},{R("O")}{mm})/SUMPRODUCT({R("R")},{R("P")}{mm}),"n/a")').number_format = "0.0%"
        RES.cell(row=rr, column=4, value=f'=IFERROR(AVERAGEIFS({R("U")},{R("E")},0,{R("I")},">=0.8"{crit}),"n/a")').number_format = P2
        RES.cell(row=rr, column=5, value=f'=COUNTIFS({R("E")},0,{R("I")},">=0.8",{R("U")},">=-10"{crit})')
        RES.cell(row=rr, column=6, value=f'=IFERROR(AVERAGEIFS({R("U")},{R("E")},1,{R("I")},">=0.8"{crit}),"n/a")').number_format = P2
        RES.cell(row=rr, column=7, value=f'=IFERROR(AVERAGEIFS({R("X")}{crit if crit else ""}),"n/a")' if crit else f'=AVERAGE({R("X")})').number_format = "0.0%"
        rr += 1
    RES.cell(row=rr, column=1, value=("Reading: Fed weeks bring macro news on top of company news. If the skill or the asymmetry differs there, "
                                      "the scenario card should flag reports that fall in a Fed week.")).font = NOTE

    # ---- Live: next Fed meeting ----
    LIV = wb["Live"]
    nxt = fm[fm["open_"] & fm["p_now"].notna()]
    if len(nxt):
        d0 = nxt["date"].min(); nk = nxt[(nxt["date"] == d0)]
        n_live = 0
        while LIV.cell(row=5 + n_live, column=1).value is not None:
            n_live += 1
        r0 = 5 + n_live + 3
        LIV.cell(row=r0, column=1, value=f"Next Fed decision: {d0.date()} (open contracts, price now)").font = SEC
        head(LIV, r0 + 1, ["Outcome", "Kalshi", "Polymarket"])
        for k, lab in enumerate(FED_LABELS):
            LIV.cell(row=r0 + 2 + k, column=1, value=lab)
            for j, v in enumerate(("Kalshi", "Polymarket")):
                hit = nk[(nk["venue"] == v) & (nk["outcome"] == lab)]
                if len(hit):
                    raw = 5 + int(hit.index[0])
                    LIV.cell(row=r0 + 2 + k, column=2 + j, value=f"=Raw_Fed_Markets!I{raw}").number_format = "0%"
        LIV.cell(row=r0 + 8, column=1, value=(f"Earnings reports within {FED_WEEK_DAYS} days of this decision fall in a Fed week: "
                                              "see Results part P for how those reports behaved in the past.")).font = NOTE


def build_kalshi(wb, km, kodds, prices, PE, EE, download_utc, head, ASM, st):
    """Kalshi KPI sheets: Raw_K_Markets, Raw_K_Odds, Kalshi_Events, Kalshi_Results. All statistics are formulas."""
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter as L
    BOLD, NOTE, TITLE, SEC, BLUE = st["BOLD"], st["NOTE"], st["TITLE"], st["SEC"], st["BLUE"]
    i0 = wb.sheetnames.index("Assumptions")
    KR = wb.create_sheet("Kalshi_Results", i0); KE = wb.create_sheet("Kalshi_Events", i0 + 1)
    RK, RKO = wb.create_sheet("Raw_K_Markets"), wb.create_sheet("Raw_K_Odds")

    # ---------- raw ----------
    RK["A1"] = (f"Kalshi KPI contracts as downloaded: {KALSHI}/historical/markets and /markets?status=settled, series tagged KPIs. "
                f"One contract per KPI event (the most traded 'above X' strike). Downloaded {download_utc} UTC."); RK["A1"].font = BOLD
    RK["A2"] = ("Stock ticker and report date are matched by the script: ticker from the series name, report date = the Yahoo "
                "earnings date within 2 days of the contract close. Cut-off = 10:00 UTC on the report day or 1 hour before close, whichever is earlier.")
    head(RK, 4, ["Market ticker", "Series", "Event", "Question", "Stock ticker", "Report date", "Strike (KPI threshold)",
                 "Actual KPI (settlement value)", "Outcome (1 = above)", "Volume (contracts)", "Close time (UTC)", "Cut-off (unix)",
                 "Report date found on Yahoo"])
    km = km.reset_index(drop=True)
    for i, x in km.iterrows():
        r = 5 + i
        vals = [x["mticker"], x["series"], x["event"], x["title"], x["ticker"], x["report_date"].to_pydatetime(), x["strike"],
                x["actual"] if pd.notna(x["actual"]) else None, int(x["y"]), float(x["volume"]),
                datetime.fromtimestamp(int(x["close_ts"]), timezone.utc).replace(tzinfo=None), int(x["cut"]),
                "yes" if x["date_matched"] else "no"]
        for j, v in enumerate(vals):
            RK.cell(row=r, column=1 + j, value=v)
        RK.cell(row=r, column=6).number_format = "yyyy-mm-dd"; RK.cell(row=r, column=11).number_format = "yyyy-mm-dd hh:mm"
        RK.cell(row=r, column=7).number_format = "#,##0.###"; RK.cell(row=r, column=8).number_format = "#,##0.###"
    RKO["A1"] = (f"Hourly YES prices (last trade, mid quote if no trade yet) as downloaded: {KALSHI} candlesticks, period 60 min, "
                 f"last {ODDS_WINDOW_DAYS} days before the cut-off."); RKO["A1"].font = BOLD
    head(RKO, 4, ["Contract row (Raw_K_Markets)", "Time (unix, UTC)", "Time (UTC)", "YES price"])
    r, span = 5, {}
    for i, (t, pts) in enumerate(kodds):
        a = r
        for ts, p in pts:
            RKO.cell(row=r, column=1, value=5 + i); RKO.cell(row=r, column=2, value=ts)
            RKO.cell(row=r, column=3, value=datetime.fromtimestamp(ts, timezone.utc).replace(tzinfo=None)).number_format = "yyyy-mm-dd hh:mm"
            RKO.cell(row=r, column=4, value=p); r += 1
        span[t] = (a, r - 1) if r > a else None
    for ws in (RK, RKO):
        ws.freeze_panes = "A5"

    # ---------- Kalshi_Events ----------
    KE["A1"] = "Kalshi_Events: one row per KPI contract with odds and prices (formulas linking to the Raw_ sheets)"; KE["A1"].font = TITLE
    KE["A2"] = ("p = YES price at the cut-off. KPI surprise = outcome minus p. Reaction window and abnormal return as on Events. "
                "EPS columns come from Yahoo for the same report. Columns Z to AD are helpers for Kalshi_Results part C "
                "(only events with all three numbers).")
    KE["A2"].alignment = Alignment(wrap_text=True); KE.merge_cells("A2:Y2"); KE.row_dimensions[2].height = 40
    head(KE, 4, ["Raw row", "Ticker", "Report date", "KPI contract", "Strike", "Actual KPI", "Outcome y (1 = above)", "Odds first row",
                 "Odds last row", "Cut-off (unix)", "p: P(above) at cut-off", "Volume (contracts)",
                 "Running base rate (earlier events)", "Brier: Kalshi", "Brier: base rate", "Both forecasts available",
                 "Stock return (log)", "Market return (log)", "Abnormal return", "KPI surprise y - p", "Actual / strike - 1",
                 "Analyst EPS surprise (%), capped at +/-100", "EPS beat (Yahoo, 1/0)", "Correct side (p>0.5 matches y)",
                 "All three available", "Helper: abnormal", "Helper: EPS surprise", "Helper: KPI surprise",
                 "Helper: abnormal net of EPS surprise", "Helper: KPI surprise net of EPS surprise"])
    good = [(5 + i, span[x["mticker"]], yf_symbol(x["ticker"])) for i, x in km.iterrows()
            if span.get(x["mticker"]) and yf_symbol(x["ticker"]) in prices.columns]
    E_ = 4 + max(1, len(good)); R = lambda c: f"Kalshi_Events!${c}$5:${c}${E_}"
    cM = L(2 + list(prices.columns).index(MARKET))
    for k, (raw, (a, b), sym) in enumerate(good):
        r = 5 + k
        cS = L(2 + list(prices.columns).index(sym))
        pre = f"MATCH(C{r}-1,Raw_Prices!$A$5:$A${PE},1)"; post = f"MATCH(C{r},Raw_Prices!$A$5:$A${PE},1)+1"
        win = f'Raw_Estimates!$A$5:$A${EE},B{r},Raw_Estimates!$B$5:$B${EE},">="&C{r}-3,Raw_Estimates!$B$5:$B${EE},"<="&C{r}+3'
        f = {"A": raw, "B": f"=Raw_K_Markets!E{raw}", "C": f"=Raw_K_Markets!F{raw}", "D": f"=Raw_K_Markets!D{raw}",
             "E": f"=Raw_K_Markets!G{raw}", "F": f'=IF(ISNUMBER(Raw_K_Markets!H{raw}),Raw_K_Markets!H{raw},"")',
             "G": f"=Raw_K_Markets!I{raw}", "H": a, "I": b, "J": f"=Raw_K_Markets!L{raw}",
             "K": f'=IFERROR(INDEX(Raw_K_Odds!$D${a}:$D${b},MATCH(J{r},Raw_K_Odds!$B${a}:$B${b},1)),"")',
             "L": f"=Raw_K_Markets!J{raw}",
             "M": f'=IF(COUNTIF($C$5:$C${E_},"<"&C{r})>={K_MIN_PRIOR_EVENTS},AVERAGEIFS($G$5:$G${E_},$C$5:$C${E_},"<"&C{r}),"")',
             "N": f'=IF(ISNUMBER(K{r}),(K{r}-G{r})^2,"")', "O": f'=IF(ISNUMBER(M{r}),(M{r}-G{r})^2,"")',
             "P": f"=IF(AND(ISNUMBER(N{r}),ISNUMBER(O{r})),1,0)",
             "Q": f'=IFERROR(LN(INDEX(Raw_Prices!${cS}$5:${cS}${PE},{post})/INDEX(Raw_Prices!${cS}$5:${cS}${PE},{pre})),"")',
             "R": f'=IFERROR(LN(INDEX(Raw_Prices!${cM}$5:${cM}${PE},{post})/INDEX(Raw_Prices!${cM}$5:${cM}${PE},{pre})),"")',
             "S": f'=IF(AND(ISNUMBER(Q{r}),ISNUMBER(R{r})),Q{r}-R{r},"")', "T": f'=IF(ISNUMBER(K{r}),G{r}-K{r},"")',
             "U": f'=IF(AND(ISNUMBER(F{r}),E{r}<>0),F{r}/E{r}-1,"")',
             "V": f'=IFERROR(MAX(-100,MIN(100,AVERAGEIFS(Raw_Estimates!$E$5:$E${EE},{win}))),"")',
             "W": f'=IFERROR(AVERAGEIFS(Raw_Estimates!$F$5:$F${EE},{win}),"")',
             "X": f'=IF(ISNUMBER(K{r}),IF((K{r}>0.5)=(G{r}=1),1,0),"")',
             "Y": f"=IF(AND(ISNUMBER(S{r}),ISNUMBER(T{r}),ISNUMBER(V{r})),1,0)",
             "Z": f'=IF(Y{r}=1,S{r},"")', "AA": f'=IF(Y{r}=1,V{r},"")', "AB": f'=IF(Y{r}=1,T{r},"")',
             "AC": f'=IF(Y{r}=1,Z{r}-(Kalshi_Results!$B$40+Kalshi_Results!$B$41*AA{r}),"")',
             "AD": f'=IF(Y{r}=1,AB{r}-(Kalshi_Results!$B$42+Kalshi_Results!$B$43*AA{r}),"")'}
        for c, v in f.items():
            KE[f"{c}{r}"] = v
        KE[f"C{r}"].number_format = "yyyy-mm-dd"; KE[f"E{r}"].number_format = "#,##0.###"; KE[f"F{r}"].number_format = "#,##0.###"
        for c in ("K", "M", "N", "O", "T", "W"):
            KE[f"{c}{r}"].number_format = "0.000"
        for c in ("Q", "R", "S", "U"):
            KE[f"{c}{r}"].number_format = "+0.00%;-0.00%"
        KE[f"L{r}"].number_format = "#,##0"
    KE.freeze_panes = "C5"; KE.column_dimensions["D"].width = 48

    # ---------- Kalshi_Results (fixed rows: Kalshi_Events refers to B40:B43) ----------
    KR["A1"] = "Kalshi KPI markets: does the crowd forecast operating numbers, and does a KPI surprise go with the stock's move?"
    KR["A1"].font = TITLE
    KR["A2"] = (f"{len(good)} settled Kalshi KPI contracts since {START_DATE} with odds and prices. Each asks whether a company reports a KPI "
                "(rides, bookings, users, headcount...) above a threshold. Brier and skill as on Results. The benchmark is the running "
                "share of 'above' outcomes among earlier KPI events.")
    KR["A2"].alignment = Alignment(wrap_text=True); KR.merge_cells("A2:H2"); KR.row_dimensions[2].height = 45
    pct, n3 = "0.0%", "0.000"
    cells = {4: ("A. Accuracy", None, None),
             5: ("Events", f"=COUNT({R('G')})", "0"),
             6: ("Share above the threshold (y = 1)", f"=AVERAGE({R('G')})", pct),
             7: ("Average Kalshi p", f"=AVERAGE({R('K')})", pct),
             8: ("Correct side (p above 0.5 and y = 1, or below and y = 0)", f"=AVERAGE({R('X')})", pct),
             9: ("Always guessing the more common outcome would be correct", "=MAX(B6,1-B6)", pct),
             10: ("Events with both forecasts", f"=SUM({R('P')})", "0"),
             11: ("Brier, Kalshi (same events)", f"=IFERROR(SUMPRODUCT({R('P')},{R('N')})/SUM({R('P')}),\"n/a\")", n3),
             12: ("Brier, running base rate", f"=IFERROR(SUMPRODUCT({R('P')},{R('O')})/SUM({R('P')}),\"n/a\")", n3),
             13: ("Skill vs base rate", '=IFERROR(1-B11/B12,"n/a")', pct),
             15: ("Calibration: when Kalshi says p, how often is the KPI above the threshold?", None, None),
             23: ("B. The stock's reaction: simple regressions of the abnormal return", None, None),
             31: ("C. Does the KPI surprise add to the EPS surprise? (events with all three numbers)", None, None),
             33: ("Events with abnormal return, EPS surprise and KPI surprise", f"=SUM({R('Y')})", "0"),
             34: ("R squared, EPS surprise alone", f'=IFERROR(RSQ({R("Z")},{R("AA")}),"n/a")', n3),
             35: ("Slope on the KPI surprise, holding the EPS surprise fixed", f'=IFERROR(SLOPE({R("AC")},{R("AD")}),"n/a")', "0.0000"),
             36: ("t statistic of that slope", f'=IFERROR(B35/SQRT(B38/(B33-3)/DEVSQ({R("AD")})),"n/a")', "0.0"),
             37: ("R squared, EPS surprise and KPI surprise together", f'=IFERROR(1-B38/DEVSQ({R("Z")}),"n/a")', n3),
             38: ("Residual sum of squares, both variables", f'=IFERROR(SUMSQ({R("AC")})-B35^2*DEVSQ({R("AD")}),"n/a")', "0.0000"),
             39: ("Added R squared from the KPI surprise", '=IFERROR(B37-B34,"n/a")', n3),
             40: ("Helper: intercept, abnormal return on EPS surprise", f'=IFERROR(INTERCEPT({R("Z")},{R("AA")}),0)', "0.0000"),
             41: ("Helper: slope, abnormal return on EPS surprise", f'=IFERROR(SLOPE({R("Z")},{R("AA")}),0)', "0.0000"),
             42: ("Helper: intercept, KPI surprise on EPS surprise", f'=IFERROR(INTERCEPT({R("AB")},{R("AA")}),0)', "0.0000"),
             43: ("Helper: slope, KPI surprise on EPS surprise", f'=IFERROR(SLOPE({R("AB")},{R("AA")}),0)', "0.0000"),
             46: ("D. Is a priced-in KPI worth less? Average abnormal return", None, None)}
    for rr, (lab, fm, nf) in cells.items():
        KR.cell(row=rr, column=1, value=lab)
        if fm is None:
            KR.cell(row=rr, column=1).font = SEC
        else:
            KR.cell(row=rr, column=2, value=fm).number_format = nf
    head(KR, 16, ["p from", "p to", "Events", "Average p", "Share above"])
    for k, (lo, hi) in enumerate([(0, .2), (.2, .4), (.4, .6), (.6, .8), (.8, 1.0001)]):
        rr = 17 + k
        KR.cell(row=rr, column=1, value=lo); KR.cell(row=rr, column=2, value=hi)
        KR.cell(row=rr, column=3, value=f'=COUNTIFS({R("K")},">="&A{rr},{R("K")},"<"&B{rr})')
        KR.cell(row=rr, column=4, value=f'=IFERROR(AVERAGEIFS({R("K")},{R("K")},">="&A{rr},{R("K")},"<"&B{rr}),"")').number_format = pct
        KR.cell(row=rr, column=5, value=f'=IFERROR(AVERAGEIFS({R("G")},{R("K")},">="&A{rr},{R("K")},"<"&B{rr}),"")').number_format = pct
    head(KR, 24, ["Regression of the abnormal return on", "Slope", "t", "R squared", "Events"])
    tt = lambda y, x: f'IFERROR(SLOPE({y},{x})/(STEYX({y},{x})/SQRT(DEVSQ({x}))),"n/a")'
    for k, (lab, x) in enumerate([("KPI outcome y (1 = above)", R("G")), ("KPI surprise y - p", R("T")),
                                  ("Analyst EPS surprise (%)", R("V")), ("EPS beat (Yahoo, 1/0)", R("W"))]):
        rr = 25 + k
        KR.cell(row=rr, column=1, value=lab)
        KR.cell(row=rr, column=2, value=f'=IFERROR(SLOPE({R("S")},{x}),"n/a")').number_format = "0.0000"
        KR.cell(row=rr, column=3, value="=" + tt(R("S"), x)).number_format = "0.0"
        KR.cell(row=rr, column=4, value=f'=IFERROR(RSQ({R("S")},{x}),"n/a")').number_format = n3
        KR.cell(row=rr, column=5, value=f"=COUNT({x})").number_format = "0"
    KR["A29"] = "Reading: the EPS rows are the same reports seen through Yahoo, so they show how much the KPI adds on its own terms."
    KR["A29"].font = NOTE
    KR["A44"] = ("Method: the slope in row 35 is the multiple-regression slope, found in two steps (both the abnormal return and the KPI "
                 "surprise are first cleaned of the EPS surprise, rows 40 to 43). Same result as LINEST on both variables.")
    KR["A44"].font = NOTE
    head(KR, 47, ["", "Kalshi p at least 0.8", "Kalshi p below 0.8", "Difference"])
    for k, (lab, yv) in enumerate((("KPI above the threshold (y = 1)", 1), ("KPI below the threshold (y = 0)", 0))):
        rr = 48 + k
        KR.cell(row=rr, column=1, value=lab)
        KR.cell(row=rr, column=2, value=f'=IFERROR(AVERAGEIFS({R("S")},{R("G")},{yv},{R("K")},">=0.8"),"n/a")').number_format = "+0.00%;-0.00%"
        KR.cell(row=rr, column=3, value=f'=IFERROR(AVERAGEIFS({R("S")},{R("G")},{yv},{R("K")},"<0.8"),"n/a")').number_format = "+0.00%;-0.00%"
        KR.cell(row=rr, column=4, value=f'=IFERROR(B{rr}-C{rr},"n/a")').number_format = "+0.00%;-0.00%"
    KR["A50"] = "Count of events in each cell:"; KR["A50"].font = NOTE
    for j, (yv, op) in enumerate(((1, ">=0.8"), (1, "<0.8"), (0, ">=0.8"), (0, "<0.8"))):
        KR.cell(row=50, column=2 + j, value=f'=COUNTIFS({R("G")},{yv},{R("K")},"{op}")')
    KR["A52"] = ("Kalshi KPI volumes are smaller than Polymarket's and the sample is short. One strike per KPI event is used "
                 "(the most traded). All results are associations, not causes.")
    KR["A52"].font = NOTE
    KR.column_dimensions["A"].width = 62
    for c in "BCDE":
        KR.column_dimensions[c].width = 16

    # ---------- Assumptions ----------
    rr = ASM.max_row + 2
    ASM.cell(row=rr, column=1, value="Kalshi (part two)").font = SEC; rr += 1
    for lab, v in [("Kalshi Trade API v2", f"{KALSHI}: /series (tag KPIs), /historical/markets, /markets?status=settled, candlesticks. "
                                           f"Downloaded {download_utc} UTC"),
                   ("Minimum volume (contracts)", K_MIN_VOLUME), ("Earlier events needed for the Kalshi base rate", K_MIN_PRIOR_EVENTS),
                   ("Contracts kept", "Settled 'above X' KPI contracts; one per KPI event, the most traded strike"),
                   ("Stock ticker", "From the Kalshi series name, kept only if its Yahoo earnings dates match the contract close dates"),
                   ("Report date", "Yahoo earnings date within 2 days of the contract close (column M on Raw_K_Markets says if found)")]:
        ASM.cell(row=rr, column=1, value=lab); c = ASM.cell(row=rr, column=2, value=v)
        if isinstance(v, (int, float)):
            c.font = BLUE
        rr += 1
    return len(good)


# =============================================================================
# 4. Run
# =============================================================================
def run():
    t0 = time.time()
    download_utc = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")
    print("1/5 Polymarket earnings contracts", flush=True)
    allm = fetch_markets()
    today = pd.Timestamp(datetime.now(timezone.utc).date())
    hist = allm[allm["resolved"] & (allm["report_date"] >= pd.Timestamp(START_DATE)) & (allm["volume"] >= MIN_VOLUME)
                & allm["ticker"].notna()].copy()
    hist = hist.sort_values("volume", ascending=False).drop_duplicates(["ticker", "report_date"]).sort_values("report_date")
    live = allm[(~allm["closed"]) & allm["ticker"].notna() & (allm["report_date"] >= today) & allm["p_now"].notna()]
    live = live.sort_values("volume", ascending=False).drop_duplicates(["ticker", "report_date"]).sort_values("report_date")
    print(f"   resolved contracts used: {len(hist)}, open contracts: {len(live)}", flush=True)

    km = pd.DataFrame()
    if USE_KALSHI:
        print("2/5 Kalshi KPI contracts", flush=True)
        try:
            km = fetch_kalshi()
        except Exception as e:
            print(f"   Kalshi download failed, continuing with Polymarket only: {e}", flush=True)
            km = pd.DataFrame()
    kcand = sorted({c for s in (km["series"].unique() if len(km) else []) for c in k_candidates(s)})

    tick = sorted(set(hist["ticker"]) | set(live["ticker"]))
    print(f"3/5 Yahoo Finance: EPS history for {len(tick)} Polymarket and {len(kcand)} Kalshi candidate tickers", flush=True)
    with ThreadPoolExecutor(THREADS) as ex:
        est_all = dict(zip(tick + kcand, ex.map(fetch_estimates, tick + kcand)))
    if len(km):
        km = k_resolve(km, est_all)
        km = km.sort_values("report_date").reset_index(drop=True) if len(km) else km
        print(f"   Kalshi KPI events matched to a stock: {len(km)}", flush=True)
    ktick = sorted(set(km["ticker"])) if len(km) else []
    est = [r for t in sorted(set(tick) | set(ktick)) for r in est_all.get(t, [])]
    prices = fetch_prices(sorted(set(tick) | set(ktick)), START_DATE)
    print(f"   price columns: {prices.shape[1]}, EPS history rows: {len(est)}", flush=True)

    print("4/5 Hourly odds before each report (Polymarket and Kalshi)", flush=True)
    with ThreadPoolExecutor(THREADS) as ex:
        pts = list(ex.map(lambda x: fetch_odds(x[0], x[1]), zip(hist["token"], hist["report_date"])))
        kpts = list(ex.map(k_odds, km.to_dict("records"))) if len(km) else []
    odds = list(zip(hist["slug"], pts))
    kodds = list(zip(km["mticker"], kpts)) if len(km) else []
    print(f"   Polymarket contracts with odds: {sum(1 for p in pts if p)}, Kalshi contracts with odds: {sum(1 for p in kpts if p)}", flush=True)

    opts, hist_px = None, None
    if USE_OPTIONS and len(live):
        try:
            hist_px = fetch_prices(sorted(set(live["ticker"])), HIST_START)
        except Exception:
            hist_px = None
        with ThreadPoolExecutor(THREADS) as ex:
            opts = list(ex.map(lambda x: fetch_option_move(x[0], x[1]), zip(live["ticker"], live["report_date"])))
        print(f"   option-implied moves for upcoming reports: {sum(1 for o in opts if o)} of {len(live)}", flush=True)

    print("   Extra free data (SEC EDGAR and FINRA)", flush=True)
    edgar, short = None, None
    dmin = str((hist["report_date"].min() - timedelta(days=5)).date()) if len(hist) else START_DATE
    allt = sorted(set(hist["ticker"]) | set(live["ticker"]))
    if USE_EDGAR:
        print("   SEC EDGAR: release times of earnings filings", flush=True)
        try:
            edgar = fetch_all_edgar(allt, dmin)
        except Exception as e:
            print(f"   EDGAR skipped: {e}", flush=True)
    if USE_FINRA:
        print("   FINRA: short interest", flush=True)
        try:
            short = fetch_short(allt, str((pd.Timestamp(dmin) - timedelta(days=40)).date()), str(datetime.now(timezone.utc).date()))
        except Exception as e:
            print(f"   FINRA skipped: {e}", flush=True)

    fed = None
    if USE_FED:
        print("   Fed meetings: Kalshi, Polymarket and FRED", flush=True)
        try:
            fed = fetch_fed()
        except Exception as e:
            print(f"   Fed module skipped: {e}", flush=True)

    print("5/5 Workbook", flush=True)
    n, nk = build(hist, odds, prices, est, live.to_dict("records"), download_utc, km, kodds, opts, hist_px, edgar, short, fed)
    import zipfile
    with zipfile.ZipFile(RAW_ZIP, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("polymarket_markets.csv", allm.to_csv(index=False))
        z.writestr("polymarket_odds.csv", pd.DataFrame([(s, t, p) for s, ps in odds for t, p in ps], columns=["slug", "ts", "p"]).to_csv(index=False))
        z.writestr("prices_yfinance.csv", prices.to_csv())
        z.writestr("eps_yfinance.csv", pd.DataFrame(est).to_csv(index=False))
        if edgar is not None and len(edgar):
            z.writestr("sec_edgar_earnings_filings.csv", edgar.to_csv(index=False))
        if fed is not None and len(fed.get("markets", [])):
            z.writestr("fed_contracts.csv", fed["markets"].to_csv(index=False))
            z.writestr("fed_odds.csv", pd.DataFrame([(s_, t, p) for s_, ps in fed["odds"] for t, p in ps], columns=["market", "ts", "p"]).to_csv(index=False))
            yfile = "yahoo_fvx_5y_yield.csv" if "Yahoo" in fed.get("yname", "") else "fred_dgs2.csv"
            z.writestr(yfile, fed["y2"].to_csv(index=False))
            if fed["target"] is not None and len(fed["target"]):
                z.writestr("fred_dfedtaru.csv", fed["target"].to_csv(index=False))
        if short is not None and len(short):
            z.writestr("finra_short_interest.csv", short.to_csv(index=False))
        if len(km):
            z.writestr("kalshi_kpi_markets.csv", km.to_csv(index=False))
            z.writestr("kalshi_odds.csv", pd.DataFrame([(s, t, p) for s, ps in kodds for t, p in ps], columns=["market", "ts", "p"]).to_csv(index=False))
    print(f"\nDone in {time.time() - t0:.0f} s: {OUTPUT_FILE} with {n} Polymarket events, {nk} Kalshi KPI events, "
          f"{len(live)} upcoming. Raw data: {RAW_ZIP}", flush=True)
    print("Open the workbook in Excel so the formulas calculate. Start with the Results sheet.", flush=True)
    out = OUTPUT_FILE
    if RECALC_IN_COLAB:
        try:
            import subprocess, os, shutil
            print("Calculating every formula (LibreOffice), about 2 minutes...", flush=True)
            if not shutil.which("soffice"):
                subprocess.run("apt-get -qq update && apt-get -qq install -y libreoffice-calc", shell=True, check=True,
                               stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", "calculated", OUTPUT_FILE],
                           check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=1500)
            calc = os.path.join("calculated", OUTPUT_FILE)
            if os.path.exists(calc):
                shutil.move(calc, OUTPUT_FILE)
                print(f"All formulas calculated and saved in {OUTPUT_FILE}", flush=True)
        except Exception as e:
            print(f"Calculation step skipped ({e}). Open the workbook in desktop Excel to calculate it.", flush=True)
    try:
        from google.colab import files
        files.download(out); files.download(RAW_ZIP)
    except Exception:
        print("If the download does not start: click the folder icon on the left, right-click the files and choose Download.", flush=True)


run()
